"""거래내역 CSV/Excel 내보내기.

UI는 이 모듈의 함수만 호출한다. 실제 SQL 조회는 transaction_service를 통해서만
하고(리포지토리를 직접 호출하지 않음), 여기서는 파일 형식으로 변환만 담당한다.
"""
import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from services import transaction_service

TYPE_LABELS = {"income": "수입", "expense": "지출"}

GENERAL_COLUMNS = [
    "거래일자", "거래시간", "수입/지출", "거래내용", "금액", "잔액",
    "카테고리", "거래처", "업무유형", "부가세", "증빙", "메모",
]
GENERAL_COLUMN_WIDTHS = [12, 10, 8, 32, 14, 14, 14, 16, 14, 8, 10, 24]

TAX_COLUMNS = [
    "거래일자", "거래내용", "입금", "출금", "거래처",
    "회계구분", "계정과목", "업무유형", "부가세 여부", "증빙", "메모",
]
TAX_COLUMN_WIDTHS = [12, 32, 14, 14, 16, 10, 14, 14, 10, 10, 24]

HEADER_FILL_COLOR = "2A78D6"


def _get_rows(filters: dict | None) -> list[dict]:
    return transaction_service.list_transactions(filters)


def _style_header(ws) -> None:
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color=HEADER_FILL_COLOR, end_color=HEADER_FILL_COLOR, fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")


def _apply_column_widths(ws, widths: list[int]) -> None:
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def export_general_csv(filters: dict | None = None) -> bytes:
    rows = _get_rows(filters)
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(GENERAL_COLUMNS)
    for t in rows:
        writer.writerow(
            [
                t["transaction_date"],
                t["transaction_time"] or "",
                TYPE_LABELS[t["transaction_type"]],
                t["description"],
                t["amount"],
                t["balance"] if t["balance"] is not None else "",
                t["category_name"] or "",
                t["client_name"] or "",
                t["work_type_name"] or "",
                t["vat_status"],
                t["evidence_status"],
                t["memo"] or "",
            ]
        )
    # 엑셀에서 한글이 깨지지 않도록 UTF-8 BOM을 붙인다.
    return ("﻿" + buf.getvalue()).encode("utf-8")


def export_general_excel(filters: dict | None = None) -> bytes:
    rows = _get_rows(filters)
    wb = Workbook()
    ws = wb.active
    ws.title = "거래내역"

    ws.append(GENERAL_COLUMNS)
    _style_header(ws)

    for t in rows:
        ws.append(
            [
                t["transaction_date"],
                t["transaction_time"] or "",
                TYPE_LABELS[t["transaction_type"]],
                t["description"],
                t["amount"],
                t["balance"],
                t["category_name"] or "",
                t["client_name"] or "",
                t["work_type_name"] or "",
                t["vat_status"],
                t["evidence_status"],
                t["memo"] or "",
            ]
        )

    # 금액(E)/잔액(F) 열은 숫자 셀로 저장하고 1,000단위 구분 서식을 적용한다.
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=6):
        for cell in row:
            if cell.value is not None:
                cell.number_format = "#,##0"

    _apply_column_widths(ws, GENERAL_COLUMN_WIDTHS)
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_tax_excel(filters: dict | None = None) -> bytes:
    """세무사 전달용 서식 (법인 통장 기준).

    입금/출금은 통장 방향(transaction_type)을 그대로 나타내며, 매출/비용
    여부는 별도의 회계구분 열로 확인한다 (입금이 곧 매출이 아니다 - 예:
    대표자 가수금 입금은 입금이지만 회계구분은 '비매출입금'일 수 있다).
    """
    rows = _get_rows(filters)
    wb = Workbook()
    ws = wb.active
    ws.title = "세무사 전달용"

    ws.append(TAX_COLUMNS)
    _style_header(ws)

    for t in rows:
        is_income = t["transaction_type"] == "income"
        ws.append(
            [
                t["transaction_date"],
                t["description"],
                t["amount"] if is_income else None,
                None if is_income else t["amount"],
                t["client_name"] or "",
                t["accounting_type"],
                t["account_name"] or "",
                t["work_type_name"] or "",
                t["vat_status"],
                t["evidence_status"],
                t["memo"] or "",
            ]
        )

    for row in ws.iter_rows(min_row=2, min_col=3, max_col=4):
        for cell in row:
            if cell.value is not None:
                cell.number_format = "#,##0"

    _apply_column_widths(ws, TAX_COLUMN_WIDTHS)
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
