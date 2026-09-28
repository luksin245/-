"""10단계(부가세 정리: 공급가액/부가세 분리, 추천, 한꺼번에 적용, 예상 부가세) 검증 스크립트.

실행 방법:
    python tests/test_stage10.py

실제 data/finance.db는 절대 건드리지 않고 tests/ 아래 임시 DB만 사용한 뒤 종료 시 삭제한다.
금액은 모두 가짜 데이터다.
"""
import io
import shutil
import sqlite3
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import db.database as database  # noqa: E402

TESTS_DIR = Path(__file__).resolve().parent
TEMP_DB_PATH = TESTS_DIR / "_stage10_test.db"
OLD_DB_PATH = TESTS_DIR / "_stage10_old_test.db"
BACKUP_DIR = TESTS_DIR / "backups"
for p in (TEMP_DB_PATH, OLD_DB_PATH):
    if p.exists():
        p.unlink()
shutil.rmtree(BACKUP_DIR, ignore_errors=True)

from openpyxl import Workbook, load_workbook  # noqa: E402

from db.database import init_db  # noqa: E402
from services import (  # noqa: E402
    card_service,
    category_rule_service,
    dashboard_service,
    export_service,
    transaction_service,
    vat_service,
)
from utils.validators import ValidationError  # noqa: E402
from utils.vat import looks_vat_included, split_vat  # noqa: E402

results: list[tuple[str, bool]] = []


def check(name: str, condition: bool) -> None:
    results.append((name, bool(condition)))
    print(f"[{'PASS' if condition else 'FAIL'}] {name}")


def expect_error(name: str, fn) -> None:
    try:
        fn()
        check(name, False)
    except ValidationError:
        check(name, True)


def columns_of(conn: sqlite3.Connection, table: str) -> set[str]:
    return {row[1] for row in conn.execute(f"PRAGMA table_info({table})")}


def check_migration() -> None:
    """부가세 컬럼이 없던 예전 DB를 열어도 데이터 손실 없이 컬럼이 추가되는지."""
    conn = sqlite3.connect(OLD_DB_PATH)
    conn.executescript(
        """
        CREATE TABLE category_rules (
            id INTEGER PRIMARY KEY AUTOINCREMENT, keyword TEXT NOT NULL,
            match_field TEXT NOT NULL DEFAULT 'description',
            suggested_category_id INTEGER, suggested_client_id INTEGER, suggested_work_type_id INTEGER,
            hit_count INTEGER NOT NULL DEFAULT 0, is_active INTEGER NOT NULL DEFAULT 1, created_at TEXT NOT NULL
        );
        CREATE TABLE ocr_raw_lines (
            id INTEGER PRIMARY KEY AUTOINCREMENT, document_id INTEGER NOT NULL, line_no INTEGER NOT NULL,
            raw_date TEXT, raw_time TEXT, raw_description TEXT, raw_income INTEGER, raw_expense INTEGER,
            raw_balance INTEGER, confidence REAL, status TEXT NOT NULL DEFAULT '확인필요',
            suggested_category_id INTEGER, suggested_client_id INTEGER, suggested_work_type_id INTEGER,
            is_confirmed INTEGER NOT NULL DEFAULT 0, linked_transaction_id INTEGER
        );
        CREATE TABLE card_statements (
            id INTEGER PRIMARY KEY AUTOINCREMENT, card_name TEXT NOT NULL, period_start TEXT NOT NULL,
            period_end TEXT NOT NULL, billed_total INTEGER NOT NULL, settlement_transaction_id INTEGER,
            memo TEXT, created_at TEXT NOT NULL
        );
        CREATE TABLE card_transactions (
            id INTEGER PRIMARY KEY AUTOINCREMENT, statement_id INTEGER NOT NULL, use_date TEXT NOT NULL,
            merchant TEXT NOT NULL, amount INTEGER NOT NULL, category_id INTEGER, account_id INTEGER,
            accounting_type TEXT NOT NULL DEFAULT '미분류', memo TEXT, created_at TEXT NOT NULL, updated_at TEXT NOT NULL
        );
        INSERT INTO card_statements VALUES (1, '옛카드', '2026-01-01', '2026-01-31', 11000, NULL, NULL, 'x');
        INSERT INTO card_transactions (statement_id, use_date, merchant, amount, created_at, updated_at)
            VALUES (1, '2026-01-05', '옛가맹점', 11000, 'x', 'x');
        """
    )
    conn.commit()
    conn.close()

    database.DB_PATH = OLD_DB_PATH
    init_db()
    conn = sqlite3.connect(OLD_DB_PATH)
    ok = (
        "vat_status" in columns_of(conn, "card_transactions")
        and "suggested_vat_status" in columns_of(conn, "category_rules")
        and "suggested_vat_status" in columns_of(conn, "ocr_raw_lines")
        and conn.execute("SELECT merchant, vat_status FROM card_transactions").fetchall() == [("옛가맹점", "불명")]
    )
    conn.close()
    check("1. 예전 DB에 부가세 컬럼을 추가해도 기존 카드 사용내역은 그대로('불명')", ok)
    OLD_DB_PATH.unlink()


def main() -> bool:
    # ---------------- 계산 ----------------
    check("2. 110,000원 -> 공급가액 100,000 + 부가세 10,000", split_vat(110_000) == (100_000, 10_000))
    check("3. 329,725원 -> 299,750 + 29,975", split_vat(329_725) == (299_750, 29_975))
    check("4. 10,000원 -> 9,091 + 909 (원 미만 반올림, 합계는 원래 금액과 같음)", split_vat(10_000) == (9_091, 909))
    check("5. 취소 -3,300원 -> -3,000 + -300", split_vat(-3_300) == (-3_000, -300))
    check("6. 금액만으로는 약한 신호: 이자 418원도 11로 나누어떨어짐 / 62,500원은 아님",
          looks_vat_included(418) and not looks_vat_included(62_500))

    # ---------------- 추천 규칙 ----------------
    s = vat_service.suggest_vat_status
    check("7. 매출·비용 + 11의 배수면 '과세' 추천", s(110_000, "매출") == {"vat_status": "과세", "reason": vat_service.REASON_AMOUNT})
    check("8. 이자(비매출입금)·가수금은 11의 배수여도 추천 안 함", s(418, "비매출입금") is None and s(2_200_000, "비매출입금") is None)
    check("9. 회계구분 미분류면 금액만으로 추천 안 함", s(55_000, "미분류") is None)
    check("10. 규칙에 부가세 여부가 있으면 그 값이 우선", s(110_000, "비용", "면세") == {"vat_status": "면세", "reason": vat_service.REASON_RULE})

    check_migration()

    database.DB_PATH = TEMP_DB_PATH
    init_db()

    # ---------------- 자동분류 규칙 ----------------
    category_rule_service.create_rule(
        keyword="CMS사용료", suggested_accounting_type="비용", suggested_vat_status="과세"
    )
    check("11. 규칙에 추천 부가세 여부 저장/추천", category_rule_service.suggest_for("FB자동 CMS사용료", None)["vat_status"] == "과세")
    expect_error("12. 잘못된 추천 부가세 값 거부",
                 lambda: category_rule_service.create_rule(keyword="x", suggested_vat_status="10%"))

    # ---------------- 통장 거래 ----------------
    def tx(day, kind, desc, amount, acct):
        return transaction_service.create_transaction(
            transaction_date=f"2026-09-{day:02d}", transaction_type=kind, description=desc, amount=amount,
            accounting_type=acct,
        )

    t_sales = tx(4, "income", "FB자금 CMS집금", 329_725, "매출")
    t_interest = tx(19, "income", "이자", 418, "비매출입금")
    t_fee = tx(7, "expense", "FB자동 CMS사용료", 110, "미분류")
    t_loan = tx(10, "income", "대표자 가수금", 2_200_000, "비매출입금")
    t_labor = tx(19, "expense", "화성동탄 노무", 62_500, "비용")
    t_unclassified = tx(15, "income", "김봉수", 55_000, "미분류")
    t_card = tx(15, "expense", "카드결 신한카드법인", 83_754, "비용")  # 아래 명세서 총합계와 같은 금액

    review = vat_service.get_bank_review("2026-07-01", "2026-09-30")
    by_id = {r["id"]: r for r in review["rows"]}
    check("13. 추천: CMS집금(매출, 11의 배수) 과세 / CMS사용료(규칙) 과세",
          by_id[t_sales]["suggested_vat_status"] == "과세" and by_id[t_fee]["suggest_reason"] == vat_service.REASON_RULE)
    check("14. 이자·가수금·11의 배수가 아닌 노무비·미분류 입금은 추천 없음",
          all(by_id[i]["suggested_vat_status"] is None for i in (t_interest, t_loan, t_labor, t_unclassified)))
    check("15. 회계구분 미분류라서 추천 못 한 건수 안내(1건)", review["unclassified_count"] == 1)
    check("16. 카드값 결제 출금 83,754원(비용)도 11의 배수라 과세로 추천됨 -> 명세서 연결로 정리해야 함",
          by_id[t_card]["suggested_vat_status"] == "과세")

    saved = vat_service.apply_bank_vat(
        [{"id": t_sales, "vat_status": "과세"}, {"id": t_fee, "vat_status": "과세"}, {"id": t_labor, "vat_status": "불명"}]
    )
    after = transaction_service.get_transaction(t_sales)
    check("17. 확인 후 저장: '불명'으로 둔 줄은 건너뛰고 2건만 저장",
          saved == 2 and transaction_service.get_transaction(t_labor)["vat_status"] == "불명")
    check("18. 부가세 여부만 바뀌고 금액·회계구분은 그대로", after["amount"] == 329_725 and after["accounting_type"] == "매출")
    expect_error("19. 잘못된 부가세 값 저장 거부", lambda: vat_service.apply_bank_vat([{"id": t_sales, "vat_status": "10%"}]))

    # ---------------- 법인카드 ----------------
    lines = [
        dict(use_date="2026-08-07", merchant="가짜다이소", amount=27_500, accounting_type="비용", vat_status="과세"),
        dict(use_date="2026-08-14", merchant="가짜택시", amount=14_600, accounting_type="비용"),
        dict(use_date="2026-08-28", merchant="가짜문구", amount=16_500, accounting_type="비용"),
        dict(use_date="2026-08-20", merchant="GAKJA SOFTWARE", amount=28_454, accounting_type="비용"),
        dict(use_date="2026-08-29", merchant="가짜다이소 취소", amount=-3_300, accounting_type="비용", vat_status="과세"),
    ]
    expect_error("20. 카드 사용내역의 잘못된 부가세 값 거부",
                 lambda: card_service.create_statement("가짜카드", "2026-08-01", "2026-08-31", 1_000,
                                                       [dict(use_date="2026-08-01", merchant="x", amount=1_000, vat_status="포함")]))
    sid = card_service.create_statement("가짜카드", "2026-08-01", "2026-08-31", 83_754, lines)
    saved_lines = card_service.get_lines(statement_id=sid)  # 이용일자 순
    check("21. 카드 사용내역에 부가세 여부 저장(지정 안 하면 '불명')",
          {l["merchant"]: l["vat_status"] for l in saved_lines}
          == {"가짜다이소": "과세", "가짜택시": "불명", "GAKJA SOFTWARE": "불명", "가짜문구": "불명", "가짜다이소 취소": "과세"})

    card_review = vat_service.get_card_review("2026-07-01", "2026-09-30")
    card_suggested = {r["merchant"]: r["suggested_vat_status"] for r in card_review["rows"]}
    check("22. 카드 추천: 16,500원(비용, 11의 배수)만 과세 추천",
          card_suggested == {"가짜택시": None, "가짜문구": "과세", "GAKJA SOFTWARE": None})
    moonggu_id = next(l["id"] for l in saved_lines if l["merchant"] == "가짜문구")
    check("23. 카드 추천값 저장", vat_service.apply_card_vat([{"id": moonggu_id, "vat_status": "과세"}]) == 1)

    card_service.link_settlement(sid, t_card, reclassify=True)
    linked = transaction_service.get_transaction(t_card)
    check("24. 카드값 결제 출금 연결 시 부가세 '해당없음'으로 (두 번 계산 방지)",
          linked["vat_status"] == "해당없음" and linked["accounting_type"] == "비비용출금")
    check("25. 연결 후에는 부가세 확인 목록에서 빠짐",
          t_card not in {r["id"] for r in vat_service.get_bank_review("2026-07-01", "2026-09-30")["rows"]})

    # ---------------- 예상 부가세 ----------------
    summary = vat_service.get_summary("2026-07-01", "2026-09-30")
    check("26. 매출세액 = 과세 입금 부가세 29,975원", summary["sales_vat"] == 29_975)
    check("27. 매입세액 = 통장 10원 + 카드 2,500+1,500-300 = 3,710원",
          summary["purchase_vat_bank"] == 10 and summary["purchase_vat_card"] == 3_700 and summary["purchase_vat"] == 3_710)
    check("28. 예상 납부세액 = 26,265원", summary["estimated_payable"] == 26_265)
    check("29. 공급가액 기준 조정용 부가세(매출 29,975 / 비용 3,700)", summary["revenue_vat"] == 29_975 and summary["cost_vat"] == 3_700)
    check("30. 아직 '불명'인 매출·비용 3건(노무비 + 카드 2건)은 계산에서 빠졌다고 안내", summary["unknown_count"] == 3)
    dash = dashboard_service.get_dashboard_data(date(2026, 7, 1), date(2026, 9, 30))
    check("31. 대시보드에도 같은 예상 부가세", dash["vat_summary"]["estimated_payable"] == 26_265)

    q = vat_service.get_period_range
    check("32. 분기 빠른 선택 (이번 분기/지난 분기/연초의 지난 분기)",
          q("이번 분기", today=date(2026, 9, 28)) == (date(2026, 7, 1), date(2026, 9, 30))
          and q("지난 분기", today=date(2026, 9, 28)) == (date(2026, 4, 1), date(2026, 6, 30))
          and q("지난 분기", today=date(2026, 2, 10)) == (date(2025, 10, 1), date(2025, 12, 31))
          and q("이번 분기", today=date(2026, 11, 3)) == (date(2026, 10, 1), date(2026, 12, 31)))

    # ---------------- 세무사 전달용 Excel ----------------
    wb = load_workbook(io.BytesIO(export_service.export_tax_excel({"start_date": "2026-08-01", "end_date": "2026-09-30"})))
    ws = wb["세무사 전달용"]
    header = [c.value for c in ws[1]]
    rows = {r[1]: dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)}
    check("33. 세무사 Excel: 과세 거래는 공급가액/부가세액 채움",
          rows["FB자금 CMS집금"]["공급가액"] == 299_750 and rows["FB자금 CMS집금"]["부가세액"] == 29_975)
    check("34. 과세가 아닌 거래는 공급가액/부가세액 비움", rows["이자"]["공급가액"] is None and rows["이자"]["부가세액"] is None)
    card_ws = wb["법인카드 사용내역"]
    card_header = [c.value for c in card_ws[1]]
    card_rows = {r[1]: dict(zip(card_header, r)) for r in card_ws.iter_rows(min_row=2, values_only=True)}
    check("35. 카드 시트에도 부가세 여부/공급가액/부가세액",
          card_rows["가짜다이소"]["부가세 여부"] == "과세" and card_rows["가짜다이소"]["공급가액"] == 25_000
          and card_rows["가짜다이소"]["부가세액"] == 2_500 and card_rows["가짜택시"]["부가세액"] is None)

    # ---------------- 카드 파일 불러오기의 부가세 열 ----------------
    def xlsx(data_rows):
        book = Workbook()
        for row in data_rows:
            book.active.append(row)
        buf = io.BytesIO()
        book.save(buf)
        return buf.getvalue()

    parsed = card_service.parse_lines_file(
        xlsx([["이용일자", "가맹점명", "청구금액", "부가세 여부"], ["2026-08-07", "가짜", 11_000, "과세"]]), "a.xlsx"
    )
    check("36. 불러오기 파일의 '부가세 여부' 열도 읽음", parsed["rows"][0]["부가세"] == "과세")
    expect_error("37. 불러오기 파일의 잘못된 부가세 값 거부",
                 lambda: card_service.parse_lines_file(
                     xlsx([["이용일자", "가맹점명", "청구금액", "부가세"], ["2026-08-07", "가짜", 11_000, "포함"]]), "a.xlsx"))

    real_db_path = TESTS_DIR.parent / "data" / "finance.db"
    check("38. 실제 data/finance.db는 생성/변경되지 않음", not real_db_path.exists())

    print()
    total = len(results)
    passed = sum(1 for _, ok in results if ok)
    print(f"결과: {passed}/{total} 통과")

    TEMP_DB_PATH.unlink(missing_ok=True)
    shutil.rmtree(BACKUP_DIR, ignore_errors=True)
    return passed == total


if __name__ == "__main__":
    success = main()
    sys.exit(0 if success else 1)
