"""9단계(카드 사용내역 엑셀/CSV 불러오기) 검증 스크립트.

실행 방법:
    python tests/test_stage9.py

파일 읽기(파싱)만 검증하며 DB에는 아무것도 저장하지 않는다. 사용하는 값은 모두 가짜 데이터다.
"""
import io
import sys
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import db.database as database  # noqa: E402

TEMP_DB_PATH = Path(__file__).resolve().parent / "_stage9_test.db"
if TEMP_DB_PATH.exists():
    TEMP_DB_PATH.unlink()
database.DB_PATH = TEMP_DB_PATH

from openpyxl import Workbook, load_workbook  # noqa: E402

from services import card_service  # noqa: E402
from utils.validators import ValidationError  # noqa: E402

results: list[tuple[str, bool]] = []


def check(name: str, condition: bool) -> None:
    results.append((name, bool(condition)))
    print(f"[{'PASS' if condition else 'FAIL'}] {name}")


def xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def expect_error(name: str, fn, must_contain: str) -> None:
    try:
        fn()
        check(name, False)
    except ValidationError as e:
        check(name, must_contain in str(e))


def main() -> bool:
    header = ["이용일자", "가맹점명", "청구금액"]

    # 1~6. AI가 만든 형태의 엑셀 (날짜 형식 섞임, 문자열 금액, 합계 줄 포함)
    data = xlsx([
        header,
        ["2026-08-07", "주식회사 가짜다이소", 27000],
        ["26.08.14", "가짜택시", "14,600"],
        [datetime(2026, 8, 28), "가짜다이소 평택점", 16700.0],
        ["2026/08/20", "GAKJA* SOFTWARE SUB", "28,457원"],
        ["2026-08-29", "구매 취소", -3000],
        [None, None, None],
        [None, "일시불", 58300],
        [None, "총합계", 83757],
    ])
    parsed = card_service.parse_lines_file(data, "aug.xlsx")
    rows = parsed["rows"]
    check("1. 사용내역 5건만 읽음", len(rows) == 5)
    check("2. 여러 날짜 형식(2026-08-07, 26.08.14, 엑셀 날짜, 2026/08/20)을 모두 날짜로",
          [r["이용일자"] for r in rows[:4]] == [date(2026, 8, 7), date(2026, 8, 14), date(2026, 8, 28), date(2026, 8, 20)])
    check("3. 금액 표기(숫자, '14,600', 16700.0, '28,457원')를 원 단위 정수로",
          [r["청구금액"] for r in rows[:4]] == [27000, 14600, 16700, 28457])
    check("4. 취소 건 음수 유지", rows[4]["청구금액"] == -3000)
    check("5. 소계·합계 줄은 건너뛰고 알려줌", parsed["skipped"] == ["일시불", "총합계"])
    check("6. 합계 = 83,757원", sum(r["청구금액"] for r in rows) == 83_757)

    # 7. 비슷한 열 이름 + 선택 열
    parsed2 = card_service.parse_lines_file(
        xlsx([["날짜", "가맹점", "금액", "회계구분", "메모"], ["2026-08-07", "가짜", 1000, "비용", "회의"]]), "x.xlsx"
    )
    r = parsed2["rows"][0]
    check("7. 비슷한 열 이름(날짜/가맹점/금액)과 선택 열(회계구분/메모)도 읽음",
          r["청구금액"] == 1000 and r["회계구분"] == "비용" and r["메모"] == "회의")

    # 8~9. CSV (UTF-8 BOM / 윈도우 엑셀 기본 cp949)
    csv_text = "이용일자,가맹점명,청구금액\n2026-08-07,가짜다이소,\"27,000\"\n"
    check("8. CSV(UTF-8) 읽기",
          card_service.parse_lines_file(csv_text.encode("utf-8-sig"), "a.csv")["rows"][0]["청구금액"] == 27000)
    check("9. CSV(한글 윈도우 cp949) 읽기",
          card_service.parse_lines_file(csv_text.encode("cp949"), "a.csv")["rows"][0]["가맹점명"] == "가짜다이소")

    # 10~16. 잘못된 파일은 추정하지 않고 몇 행인지 알려주며 거부
    expect_error("10. 필수 열 없음 거부", lambda: card_service.parse_lines_file(xlsx([["날짜", "가맹점"], ["2026-08-07", "x"]]), "a.xlsx"), "청구금액")
    expect_error("11. 읽을 수 없는 날짜 거부(행 번호 안내)",
                 lambda: card_service.parse_lines_file(xlsx([header, ["2026-08-07", "a", 1], ["8월 둘째주", "b", 2]]), "a.xlsx"), "3행")
    expect_error("12. 원 단위가 아닌 금액(해외 이용금액 20.20 등) 거부",
                 lambda: card_service.parse_lines_file(xlsx([header, ["2026-08-20", "해외", 20.20]]), "a.xlsx"), "원 단위")
    expect_error("13. 숫자가 아닌 금액 거부",
                 lambda: card_service.parse_lines_file(xlsx([header, ["2026-08-20", "x", "USD 20"]]), "a.xlsx"), "청구금액")
    expect_error("14. 날짜가 빈 사용내역 줄 거부",
                 lambda: card_service.parse_lines_file(xlsx([header, [None, "가짜다이소", 1000]]), "a.xlsx"), "이용일자가 비어")
    expect_error("15. 허용되지 않은 회계구분 거부",
                 lambda: card_service.parse_lines_file(xlsx([header + ["회계구분"], ["2026-08-07", "x", 1, "경비"]]), "a.xlsx"), "회계구분")
    expect_error("16. 지원하지 않는 파일 형식 거부", lambda: card_service.parse_lines_file(b"x", "a.png"), "엑셀")

    # 17~18. 빈 양식
    template = card_service.build_import_template()
    wb = load_workbook(io.BytesIO(template))
    check("17. 빈 양식 첫 시트 제목 줄",
          [c.value for c in wb.active[1]][:3] == ["이용일자", "가맹점명", "청구금액"] and "작성방법" in wb.sheetnames)
    expect_error("18. 빈 양식을 그대로 올리면 '불러올 내역 없음' 안내",
                 lambda: card_service.parse_lines_file(template, "양식.xlsx"), "불러올 사용내역이 없습니다")

    real_db_path = Path(__file__).resolve().parent.parent / "data" / "finance.db"
    check("19. 실제 data/finance.db는 생성/변경되지 않음", not real_db_path.exists())

    print()
    total = len(results)
    passed = sum(1 for _, ok in results if ok)
    print(f"결과: {passed}/{total} 통과")
    TEMP_DB_PATH.unlink(missing_ok=True)
    return passed == total


if __name__ == "__main__":
    success = main()
    sys.exit(0 if success else 1)
