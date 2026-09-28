"""5단계(백업/복구, 내보내기, 통장 가져오기 OCR, 검토, 중복감지, 자동분류 규칙) 검증 스크립트.

실행 방법:
    python tests/test_stage5.py

1~4단계 테스트와 동일한 원칙을 따른다: 실제 data/finance.db는 절대 건드리지 않고,
tests/ 폴더 아래 임시 DB 파일(_stage5_test.db)과 그에 딸린 backups/temp 폴더만
사용한 뒤 종료 시 모두 삭제한다.
"""
import io
import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import db.database as database  # noqa: E402

TEMP_DB_PATH = Path(__file__).resolve().parent / "_stage5_test.db"
if TEMP_DB_PATH.exists():
    TEMP_DB_PATH.unlink()
shutil.rmtree(TEMP_DB_PATH.parent / "backups", ignore_errors=True)
shutil.rmtree(TEMP_DB_PATH.parent / "temp", ignore_errors=True)
database.DB_PATH = TEMP_DB_PATH

from db.database import init_db  # noqa: E402
from db import backup_repository  # noqa: E402
from services import (  # noqa: E402
    account_service,
    backup_service,
    category_rule_service,
    category_service,
    dedup_service,
    export_service,
    import_service,
    transaction_service,
    work_type_service,
)
from services.ocr.base import STATUS_NEEDS_REVIEW, STATUS_OK, STATUS_SUSPECTED_DUPLICATE  # noqa: E402
from utils.validators import ValidationError  # noqa: E402

FIXTURES_DIR = Path(__file__).resolve().parent / "fixtures"

results: list[tuple[str, bool]] = []


def check(name: str, condition: bool) -> None:
    results.append((name, bool(condition)))
    print(f"[{'PASS' if condition else 'FAIL'}] {name}")


def main() -> bool:
    init_db()

    income_cats = category_service.get_income_categories()
    expense_cats = category_service.get_expense_categories()
    work_types = work_type_service.get_work_types()
    accounts = account_service.get_accounts()

    # =================================================================
    # A. 백업 / 복구
    # =================================================================
    tx_a = transaction_service.create_transaction(
        transaction_date="2026-09-01",
        transaction_type="income",
        description="백업전 거래A",
        amount=100_000,
        category_id=income_cats[0]["id"],
    )

    backup_info = backup_service.create_backup()
    backup_path = backup_repository.get_backup_dir() / backup_info["filename"]
    check("1. 백업 파일이 실제로 생성됨", backup_path.exists())

    backups_after_one = backup_service.list_backups()
    check("2. list_backups에 방금 만든 백업이 포함됨", any(b["filename"] == backup_info["filename"] for b in backups_after_one))

    db_info = backup_service.get_db_info()
    check("3. get_db_info의 db_path가 현재 DB 경로와 일치", Path(db_info["db_path"]) == TEMP_DB_PATH)
    check("4. get_db_info의 last_backup_time이 채워짐", db_info["last_backup_time"] is not None)

    try:
        backup_service.restore_backup("존재하지않는파일.db")
        check("5. 존재하지 않는 백업 복구 시도 거부", False)
    except ValidationError:
        check("5. 존재하지 않는 백업 복구 시도 거부", True)

    tx_b = transaction_service.create_transaction(
        transaction_date="2026-09-02",
        transaction_type="expense",
        description="백업후 거래B",
        amount=50_000,
        category_id=expense_cats[0]["id"],
    )
    count_before_restore = transaction_service.count_transactions()

    restore_result = backup_service.restore_backup(backup_info["filename"])

    tx_ids_after_restore = {t["id"] for t in transaction_service.list_transactions()}
    check(
        "6. 복구 후 백업 시점 이후 거래(B)는 사라지고 이전 거래(A)는 유지됨",
        tx_a in tx_ids_after_restore and tx_b not in tx_ids_after_restore,
    )
    check("6-보조. 복구 전후 거래건수가 실제로 달라짐(복구가 실제로 적용됨)", count_before_restore != transaction_service.count_transactions())

    backups_after_restore = backup_service.list_backups()
    check(
        "7. 복구 시 복구 직전 상태의 안전 백업(pre_restore)이 자동 생성됨",
        any(b["filename"] == restore_result["pre_restore_backup"] for b in backups_after_restore),
    )

    check(
        "8. 복구 후에도 마이그레이션이 재적용되어 계정과목 기본데이터가 존재함",
        len(account_service.get_accounts()) > 0,
    )

    # (이후 시나리오를 위해 거래 A 하나만 남은 깨끗한 상태를 유지)

    # =================================================================
    # B. CSV / Excel 내보내기
    # =================================================================
    transaction_service.create_transaction(
        transaction_date="2026-09-03",
        transaction_type="expense",
        description="내보내기용 거래",
        amount=30_000,
        category_id=expense_cats[0]["id"],
        accounting_type="비용",
        account_id=accounts[0]["id"] if accounts else None,
    )

    csv_bytes = export_service.export_general_csv()
    csv_text = csv_bytes.decode("utf-8-sig")
    csv_line_count = len([l for l in csv_text.splitlines() if l.strip()])
    check("9. CSV 내보내기 행 개수(헤더+거래건수) 일치", csv_line_count == 1 + transaction_service.count_transactions())
    check("10. CSV 내보내기에 거래 설명 포함", "내보내기용 거래" in csv_text)

    from openpyxl import load_workbook

    general_xlsx = export_service.export_general_excel()
    wb = load_workbook(io.BytesIO(general_xlsx))
    ws = wb.active
    header_row = [c.value for c in ws[1]]
    check("11. 일반 Excel 헤더가 정의된 컬럼과 일치", header_row == export_service.GENERAL_COLUMNS)
    amount_cell = ws.cell(row=2, column=5)
    check("12. 일반 Excel 금액 셀 숫자서식(#,##0) 적용", amount_cell.number_format == "#,##0")

    tax_xlsx = export_service.export_tax_excel()
    wb_tax = load_workbook(io.BytesIO(tax_xlsx))
    ws_tax = wb_tax.active
    tax_header = [c.value for c in ws_tax[1]]
    check(
        "13. 세무사 전달용 Excel 헤더가 스펙과 정확히 일치",
        tax_header == ["거래일자", "거래내용", "입금", "출금", "거래처", "회계구분", "계정과목", "업무유형", "부가세 여부", "증빙", "메모"],
    )

    expense_row_found = False
    for row in ws_tax.iter_rows(min_row=2, values_only=True):
        if row[1] == "내보내기용 거래":
            expense_row_found = row[2] is None and row[3] == 30_000
            break
    check("14. 세무사 전달용 Excel에서 지출 거래는 출금란에만 금액이 표시됨", expense_row_found)

    filtered_csv = export_service.export_general_csv({"transaction_type": "expense"})
    filtered_text = filtered_csv.decode("utf-8-sig")
    check("15. 필터 적용된 내보내기는 조건에 맞는 행만 포함", "내보내기용 거래" in filtered_text and "백업전 거래A" not in filtered_text)

    # =================================================================
    # C. 통장 가져오기 업로드 / OCR
    # =================================================================
    try:
        import_service.process_uploaded_file(b"dummy", "statement.txt")
        check("16. 지원하지 않는 파일 형식 업로드 거부", False)
    except ValidationError:
        check("16. 지원하지 않는 파일 형식 업로드 거부", True)

    temp_dir = TEMP_DB_PATH.parent / "temp"
    with open(FIXTURES_DIR / "fake_bank_statement.png", "rb") as f:
        png_bytes = f.read()

    doc_id = import_service.process_uploaded_file(png_bytes, "fake_bank_statement.png")
    check(
        "17. OCR 처리 직후 업로드 원본 임시파일이 남지 않음(개인정보 보호)",
        not any(temp_dir.glob("*")) if temp_dir.exists() else True,
    )

    lines = import_service.get_review_lines(doc_id)
    check("18. OCR로 거래 후보가 1개 이상 추출됨", len(lines) > 0)
    check(
        "19. 어떤 후보도 입금/출금이 동시에 채워지지 않음(추정 금지 원칙)",
        all(not (l["raw_income"] is not None and l["raw_expense"] is not None) for l in lines),
    )
    check(
        "20. '정상' 상태 후보는 반드시 날짜와 입금/출금 중 하나를 가짐",
        all(
            (l["raw_date"] is not None and (l["raw_income"] is not None) != (l["raw_expense"] is not None))
            for l in lines
            if l["status"] == STATUS_OK
        ),
    )
    check("21. 검토 후보가 line_no 순서대로 정렬됨", [l["line_no"] for l in lines] == sorted(l["line_no"] for l in lines))

    documents = import_service.list_documents()
    check("22. 업로드한 문서가 목록에 파일명과 함께 조회됨", any(d["id"] == doc_id and d["file_name"] == "fake_bank_statement.png" for d in documents))

    # 인터넷뱅킹에서 내려받는 "계좌별거래내역" 표 형식(입금액/출금액/잔액 컬럼이 있는
    # 표) PDF는 문장형 줄 단위 파서로는 인식할 수 없어 전용 표 파서를 사용한다.
    with open(FIXTURES_DIR / "fake_bank_table_statement.pdf", "rb") as f:
        table_pdf_bytes = f.read()
    table_doc_id = import_service.process_uploaded_file(table_pdf_bytes, "fake_bank_table_statement.pdf")
    table_lines = import_service.get_review_lines(table_doc_id)
    check("22-표형식. 표 형식 PDF에서 모든 행이 인식됨", len(table_lines) == 6)
    ok_table_lines = [l for l in table_lines if l["status"] == STATUS_OK]
    check("22-표형식. 정상 인식된 행은 입금/출금 중 정확히 하나만 있음",
          all((l["raw_income"] is not None) != (l["raw_expense"] is not None) for l in ok_table_lines))
    check("22-표형식. 입금/출금이 모두 없는 행(신규 등)은 확인필요로 처리됨",
          any(l["status"] == STATUS_NEEDS_REVIEW and l["raw_income"] is None and l["raw_expense"] is None for l in table_lines))
    first_row = next(l for l in table_lines if l["line_no"] == 1)
    check("22-표형식. 날짜/시간(줄바꿈 포함)이 정확히 복원됨", first_row["raw_date"] == "2026-09-01" and first_row["raw_time"] == "09:12:00")
    check("22-표형식. 입금액이 정확히 인식됨", first_row["raw_income"] == 1_500_000)
    check("22-표형식. 잔액이 정확히 인식됨", first_row["raw_balance"] == 5_000_000)
    import_service.discard_document(table_doc_id)

    # =================================================================
    # D. OCR 검토 / 확정 / 취소
    # =================================================================
    ok_line = next(l for l in lines if l["status"] == STATUS_OK)
    try:
        import_service.confirm_line(
            ok_line["id"], 1000, 500, {"transaction_date": ok_line["raw_date"], "description": "x"}
        )
        check("23. 입금/출금 동시 입력 시 확정 거부", False)
    except ValidationError:
        check("23. 입금/출금 동시 입력 시 확정 거부", True)

    try:
        import_service.confirm_line(
            ok_line["id"], None, None, {"transaction_date": ok_line["raw_date"], "description": "x"}
        )
        check("24. 입금/출금 모두 미입력 시 확정 거부", False)
    except ValidationError:
        check("24. 입금/출금 모두 미입력 시 확정 거부", True)

    count_before_confirm = transaction_service.count_transactions()
    new_tx_id = import_service.confirm_line(
        ok_line["id"],
        ok_line["raw_income"],
        ok_line["raw_expense"],
        {
            "transaction_date": ok_line["raw_date"],
            "transaction_time": ok_line["raw_time"],
            "description": ok_line["raw_description"] or "OCR거래",
            "balance": ok_line["raw_balance"],
        },
    )
    created_tx = transaction_service.get_transaction(new_tx_id)
    check("25. 확정된 후보로부터 source_type='ocr'인 실제 거래가 생성됨", created_tx is not None and created_tx["source_type"] == "ocr")
    check("25-보조. 확정 후 거래건수가 1건 증가", transaction_service.count_transactions() == count_before_confirm + 1)

    remaining_lines = import_service.get_review_lines(doc_id)
    check("26. 확정된 후보는 is_confirmed=1로 표시됨", next(l for l in remaining_lines if l["id"] == ok_line["id"])["is_confirmed"] == 1)

    unresolved_before_discard = [l for l in remaining_lines if not l["is_confirmed"]]
    count_before_discard_tx = transaction_service.count_transactions()
    discard_target = unresolved_before_discard[0]
    import_service.discard_line(discard_target["id"])
    check(
        "27. 취소(제외)한 후보는 삭제되고 새 거래도 생성되지 않음",
        discard_target["id"] not in {l["id"] for l in import_service.get_review_lines(doc_id)}
        and transaction_service.count_transactions() == count_before_discard_tx,
    )

    remaining_after_discard = [l for l in import_service.get_review_lines(doc_id) if not l["is_confirmed"]]
    for l in remaining_after_discard:
        import_service.discard_line(l["id"])
    check("28. 모든 후보가 처리되면 문서가 목록에서 자동 정리됨", all(d["id"] != doc_id for d in import_service.list_documents()))

    with open(FIXTURES_DIR / "fake_bank_statement.jpg", "rb") as f:
        jpg_bytes = f.read()
    doc_id2 = import_service.process_uploaded_file(jpg_bytes, "fake_bank_statement2.jpg")
    import_service.discard_document(doc_id2)
    check("29. 문서 전체 취소 시 즉시 목록에서 사라짐", all(d["id"] != doc_id2 for d in import_service.list_documents()))

    # =================================================================
    # E. 중복 감지
    # =================================================================
    transaction_service.create_transaction(
        transaction_date="2026-09-10",
        transaction_type="income",
        description="중복검사대상",
        amount=777_000,
        transaction_time="10:00:00",
        balance=1_000_000,
        category_id=income_cats[0]["id"],
    )
    check("30. 후보 없을 때 중복 감지 결과는 없음", dedup_service.check_duplicate({"date": "2099-01-01", "income": 1, "expense": None, "balance": None, "time": None, "description": None}) is None)

    exact_candidate = {
        "date": "2026-09-10", "time": "10:00:00", "description": "중복검사대상",
        "income": 777_000, "expense": None, "balance": 1_000_000,
    }
    check("31. 날짜+시간+내용+구분+금액+잔액이 모두 일치하면 강한 중복", dedup_service.check_duplicate(exact_candidate) == dedup_service.MATCH_STRONG)

    partial_candidate = {
        "date": "2026-09-10", "time": "11:30:00", "description": "다른내용",
        "income": 777_000, "expense": None, "balance": 1_000_000,
    }
    check("32. 날짜+금액+잔액만 일치하면 부분 중복", dedup_service.check_duplicate(partial_candidate) == dedup_service.MATCH_PARTIAL)

    diff_balance_candidate = {
        "date": "2026-09-10", "time": "10:00:00", "description": "중복검사대상",
        "income": 777_000, "expense": None, "balance": 999_999,
    }
    check("33. 날짜+금액이 같아도 잔액이 다르면 중복 아님", dedup_service.check_duplicate(diff_balance_candidate) is None)

    # OCR 파이프라인에 실제로 연결되는지: 위 거래와 완전히 동일한 내용을 담은 이미지를 만들어 업로드하면
    # '중복의심'으로 표시되어야 한다 (기존 fixture 데이터 자체가 이미 저장돼 있는 상태를 재현).
    with open(FIXTURES_DIR / "fake_bank_statement.png", "rb") as f:
        png_bytes_2 = f.read()
    # 이미 한 줄(1번, A회사 1,500,000)을 위에서 확정 저장했으므로, 같은 파일을 다시 업로드하면
    # 해당 줄이 이번에는 '중복의심'으로 표시되어야 한다.
    doc_id3 = import_service.process_uploaded_file(png_bytes_2, "fake_bank_statement3.png")
    lines3 = import_service.get_review_lines(doc_id3)
    check(
        "34. 이미 저장된 거래와 동일한 내용을 재업로드하면 '중복의심'으로 표시됨",
        any(l["status"] == STATUS_SUSPECTED_DUPLICATE for l in lines3),
    )
    check(
        "35. '중복의심' 후보는 자동으로 저장되지 않음(is_confirmed=0 유지)",
        all(not l["is_confirmed"] for l in lines3 if l["status"] == STATUS_SUSPECTED_DUPLICATE),
    )
    import_service.discard_document(doc_id3)

    doc_id4 = import_service.process_uploaded_file(png_bytes_2, "fake_bank_statement4.png")
    lines4 = import_service.get_review_lines(doc_id4)
    check(
        "36. 동일 파일을 여러 번 업로드해도 매번 동일하게 중복 판정됨(일관성)",
        any(l["status"] == STATUS_SUSPECTED_DUPLICATE for l in lines4),
    )
    import_service.discard_document(doc_id4)

    # =================================================================
    # F. 자동분류 규칙
    # =================================================================
    rule_id = category_rule_service.create_rule(
        keyword="테스트키워드",
        match_field="description",
        suggested_category_id=expense_cats[0]["id"],
        suggested_work_type_id=work_types[0]["id"],
        suggested_accounting_type="비용",
    )
    try:
        category_rule_service.create_rule(keyword="테스트키워드", match_field="description")
        check("37. 동일 키워드+매칭대상 중복 등록 거부", False)
    except ValidationError:
        check("37. 동일 키워드+매칭대상 중복 등록 거부", True)

    try:
        category_rule_service.create_rule(keyword="   ", match_field="description")
        check("38. 빈 키워드 등록 거부", False)
    except ValidationError:
        check("38. 빈 키워드 등록 거부", True)

    suggestion = category_rule_service.suggest_for("테스트키워드 관련 지출", None)
    check("39. 키워드가 거래내용에 포함되면 추천값이 반환됨(대소문자 무관 부분일치)", suggestion is not None and suggestion["accounting_type"] == "비용")
    check("40. 매칭되는 규칙이 없으면 추천값 없음", category_rule_service.suggest_for("전혀 다른 내용", None) is None)

    category_rule_service.deactivate_rule(rule_id)
    check("41. 비활성화된 규칙은 더 이상 추천에 사용되지 않음", category_rule_service.suggest_for("테스트키워드 관련 지출", None) is None)
    category_rule_service.activate_rule(rule_id)
    check("42. 재활성화하면 다시 추천에 사용됨", category_rule_service.suggest_for("테스트키워드 관련 지출", None) is not None)

    long_rule_id = category_rule_service.create_rule(keyword="테스트키워드 관련", match_field="description")
    specific_suggestion = category_rule_service.suggest_for("테스트키워드 관련 지출", None)
    check("43. 여러 규칙이 매칭되면 더 구체적인(긴) 키워드 규칙이 우선함", specific_suggestion["rule_id"] == long_rule_id)
    category_rule_service.deactivate_rule(long_rule_id)

    with open(FIXTURES_DIR / "fake_bank_statement.pdf", "rb") as f:
        pdf_bytes = f.read()
    rule_for_pdf = category_rule_service.create_rule(
        keyword="임차료", match_field="description", suggested_accounting_type="비용"
    )
    doc_id5 = import_service.process_uploaded_file(pdf_bytes, "fake_bank_statement5.pdf")
    lines5 = import_service.get_review_lines(doc_id5)
    check(
        "44. OCR 검토 후보에 자동분류 규칙의 추천값이 실제로 반영됨",
        any(l["suggested_accounting_type"] == "비용" for l in lines5),
    )
    import_service.discard_document(doc_id5)
    category_rule_service.deactivate_rule(rule_for_pdf)

    hit_count_before = category_rule_service.get_rule(rule_id)["hit_count"]
    category_rule_service.record_suggestion_used(rule_id)
    check("45. 추천 사용 시 hit_count가 증가함", category_rule_service.get_rule(rule_id)["hit_count"] == hit_count_before + 1)

    # =================================================================
    # G. 회귀 / 통합
    # =================================================================
    try:
        transaction_service.create_transaction(
            transaction_date="2026-09-11", transaction_type="income", description="잘못된회계구분",
            amount=1000, accounting_type="존재하지않는구분",
        )
        check("46. 잘못된 회계구분 값은 거부됨", False)
    except ValidationError:
        check("46. 잘못된 회계구분 값은 거부됨", True)

    regression_account_id = account_service.create_account("회귀테스트계정과목")
    regression_tx_id = transaction_service.create_transaction(
        transaction_date="2026-09-12", transaction_type="expense", description="계정과목회귀",
        amount=5000, account_id=regression_account_id,
    )
    account_service.deactivate_account(regression_account_id)
    try:
        transaction_service.update_transaction(
            transaction_id=regression_tx_id, transaction_date="2026-09-12", transaction_type="expense",
            description="계정과목회귀(메모수정)", amount=5000, account_id=regression_account_id,
        )
        check("47. 비활성화된 계정과목이어도 기존 거래에서 값을 그대로 유지하는 수정은 허용됨", True)
    except ValidationError:
        check("47. 비활성화된 계정과목이어도 기존 거래에서 값을 그대로 유지하는 수정은 허용됨", False)

    transaction_service.create_transaction(
        transaction_date="2026-09-13", transaction_type="income", description="대표자 가수금",
        amount=2_000_000, accounting_type="비매출입금",
    )
    accounting_summary = transaction_service.get_accounting_summary({"start_date": "2026-09-13", "end_date": "2026-09-13"})
    check(
        "48. 회계구분이 '비매출입금'인 입금은 총매출에 집계되지 않음(총수입과 총매출은 다른 개념)",
        accounting_summary["total_revenue"] == 0,
    )

    import db.migrations as migrations
    conn = database.get_connection()
    try:
        migrations.run_migrations(conn)
        migrations.run_migrations(conn)
        check("49. 마이그레이션을 반복 실행해도 오류 없이 안전함(멱등성)", True)
    except Exception:
        check("49. 마이그레이션을 반복 실행해도 오류 없이 안전함(멱등성)", False)
    finally:
        conn.close()

    real_db_path = Path(__file__).resolve().parent.parent / "data" / "finance.db"
    check(
        "50. 이 테스트 실행 동안 실제 data/finance.db 파일은 생성/변경되지 않음",
        database.DB_PATH == TEMP_DB_PATH and not real_db_path.exists(),
    )

    print()
    total = len(results)
    passed = sum(1 for _, ok in results if ok)
    print(f"결과: {passed}/{total} 통과")

    TEMP_DB_PATH.unlink(missing_ok=True)
    shutil.rmtree(TEMP_DB_PATH.parent / "backups", ignore_errors=True)
    shutil.rmtree(TEMP_DB_PATH.parent / "temp", ignore_errors=True)
    return passed == total


if __name__ == "__main__":
    success = main()
    sys.exit(0 if success else 1)
