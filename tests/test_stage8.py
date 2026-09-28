"""8단계(법인카드 명세서 / 사용내역 / 통장 카드값 결제 맞춰보기) 검증 스크립트.

실행 방법:
    python tests/test_stage8.py

실제 data/finance.db는 절대 건드리지 않고 tests/ 아래 임시 DB(_stage8_test.db)와
그에 딸린 backups 폴더만 사용한 뒤 종료 시 삭제한다. 금액은 사용자가 보여준 8월 명세서
구조(국내 3건 + 해외 1건, 총합계 86,757원)를 본뜬 가짜 데이터다.
"""
import io
import shutil
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import db.database as database  # noqa: E402

TEMP_DB_PATH = Path(__file__).resolve().parent / "_stage8_test.db"
BACKUP_DIR = TEMP_DB_PATH.parent / "backups"
if TEMP_DB_PATH.exists():
    TEMP_DB_PATH.unlink()
shutil.rmtree(BACKUP_DIR, ignore_errors=True)
database.DB_PATH = TEMP_DB_PATH

from openpyxl import load_workbook  # noqa: E402

from db.database import init_db  # noqa: E402
from services import (  # noqa: E402
    account_service,
    card_service,
    category_rule_service,
    category_service,
    dashboard_service,
    export_service,
    transaction_service,
)
from utils.validators import ValidationError  # noqa: E402

results: list[tuple[str, bool]] = []


def check(name: str, condition: bool) -> None:
    results.append((name, bool(condition)))
    print(f"[{'PASS' if condition else 'FAIL'}] {name}")


def expect_error(name: str, fn) -> None:
    try:
        fn()
        check(name, False)
    except ValidationError:
        check(name, True)


AUG_LINES = [
    dict(use_date="2026-08-07", merchant="가짜다이소 강남점", amount=27_000),
    dict(use_date="2026-08-14", merchant="가짜택시", amount=14_600),
    dict(use_date="2026-08-28", merchant="가짜다이소 평택점", amount=16_700),
    dict(use_date="2026-08-20", merchant="GAKJA* SOFTWARE SUB", amount=28_457),
]
AUG_TOTAL = 86_757


def main() -> bool:
    init_db()
    accounts = {a["name"]: a["id"] for a in account_service.get_accounts()}
    expense_cats = {c["name"]: c["id"] for c in category_service.get_expense_categories()}

    # 예전 규칙대로 카드결 출금을 '비용'으로 잡아둔 상태에서 시작 (연결 후 빠지는지 검증하기 위해)
    bank_settle_id = transaction_service.create_transaction(
        transaction_date="2026-09-15", transaction_type="expense", description="카드결 신한카드법인", amount=AUG_TOTAL,
        accounting_type="비용",
    )
    transaction_service.create_transaction(  # 금액만 다른 출금 (후보가 되면 안 됨)
        transaction_date="2026-09-16", transaction_type="expense", description="다른 출금", amount=AUG_TOTAL + 1,
    )
    transaction_service.create_transaction(  # 금액은 같지만 입금 (후보가 되면 안 됨)
        transaction_date="2026-09-17", transaction_type="income", description="같은 금액 입금", amount=AUG_TOTAL,
    )
    transaction_service.create_transaction(  # 금액은 같지만 60일 창 밖 (후보가 되면 안 됨)
        transaction_date="2026-12-31", transaction_type="expense", description="너무 늦은 출금", amount=AUG_TOTAL,
    )

    # ---------------- 등록 검증 ----------------
    expect_error("1. 합계가 명세서 총합계와 다르면 저장 거부(OCR 오독 86,787 사례)",
                 lambda: card_service.create_statement("신한카드 법인", "2026-08-01", "2026-08-31", 86_787, AUG_LINES))
    expect_error("2. 사용내역이 없으면 거부",
                 lambda: card_service.create_statement("신한카드 법인", "2026-08-01", "2026-08-31", 0, []))
    expect_error("3. 이용기간 종료일이 시작일보다 빠르면 거부",
                 lambda: card_service.create_statement("신한카드 법인", "2026-08-31", "2026-08-01", AUG_TOTAL, AUG_LINES))
    expect_error("4. 가맹점명 없는 줄 거부",
                 lambda: card_service.create_statement("신한카드 법인", "2026-08-01", "2026-08-31", 100,
                                                       [dict(use_date="2026-08-01", merchant=" ", amount=100)]))
    expect_error("5. 0원 줄 거부",
                 lambda: card_service.create_statement("신한카드 법인", "2026-08-01", "2026-08-31", 0,
                                                       [dict(use_date="2026-08-01", merchant="x", amount=0)]))
    expect_error("6. 잘못된 날짜 거부",
                 lambda: card_service.create_statement("신한카드 법인", "2026-08-01", "2026-08-31", 100,
                                                       [dict(use_date="2026-13-40", merchant="x", amount=100)]))
    expect_error("7. 수입 카테고리 지정 거부",
                 lambda: card_service.create_statement(
                     "신한카드 법인", "2026-08-01", "2026-08-31", 100,
                     [dict(use_date="2026-08-01", merchant="x", amount=100,
                           category_id=category_service.get_income_categories()[0]["id"])]))

    sid = card_service.create_statement("신한카드 법인", "2026-08-01", "2026-08-31", AUG_TOTAL, AUG_LINES)
    stmt = next(s for s in card_service.list_statements() if s["id"] == sid)
    check("8. 명세서 저장: 4건, 합계 86,757원", stmt["line_count"] == 4 and stmt["line_total"] == AUG_TOTAL)
    check("9. 사용내역 기본 회계구분은 '미분류'(추정하지 않음)",
          all(l["accounting_type"] == "미분류" for l in card_service.get_lines(statement_id=sid)))
    expect_error("10. 같은 카드·같은 이용기간 명세서 중복 등록 거부",
                 lambda: card_service.create_statement("신한카드 법인", "2026-08-01", "2026-08-31", AUG_TOTAL, AUG_LINES))

    refund_sid = card_service.create_statement(
        "신한카드 법인", "2026-07-01", "2026-07-31", 5_000,
        [dict(use_date="2026-07-02", merchant="구매", amount=8_000), dict(use_date="2026-07-03", merchant="구매 취소", amount=-3_000)],
    )
    check("11. 취소·환불(음수) 포함 명세서도 합계가 맞으면 저장", refund_sid is not None)

    # ---------------- 카드 사용 건이 통장 기준 합계에 섞이지 않음 ----------------
    summary = transaction_service.get_summary({"start_date": "2026-08-01", "end_date": "2026-09-30"})
    check("12. 카드 사용내역은 통장 기준 총지출에 섞이지 않음",
          summary["total_expense"] == AUG_TOTAL + AUG_TOTAL + 1)

    # ---------------- 분류 ----------------
    lines = card_service.get_lines(statement_id=sid)
    updates = [
        {"id": l["id"], "category_id": expense_cats.get("사무용품비"), "account_id": accounts.get("소모품비"),
         "accounting_type": "비용", "memo": None}
        for l in lines if l["amount"] != 14_600
    ] + [
        {"id": l["id"], "category_id": expense_cats.get("교통비"), "account_id": accounts.get("여비교통비"),
         "accounting_type": "비용", "memo": "외근"}
        for l in lines if l["amount"] == 14_600
    ]
    card_service.update_classifications(sid, updates)
    taxi = next(l for l in card_service.get_lines(statement_id=sid) if l["amount"] == 14_600)
    check("13. 사용내역 분류 저장", taxi["account_name"] == "여비교통비" and taxi["memo"] == "외근")
    expect_error("14. 다른 명세서의 사용내역은 수정 거부",
                 lambda: card_service.update_classifications(refund_sid, [dict(updates[0])]))

    category_rule_service.create_rule(keyword="택시", match_field="description",
                                      suggested_account_id=accounts.get("여비교통비"), suggested_accounting_type="비용")
    suggestion = card_service.suggest_for_merchant("티머니 택시-서울")
    check("15. 가맹점명으로 자동분류 규칙 추천", suggestion.get("account_id") == accounts.get("여비교통비")
          and suggestion.get("accounting_type") == "비용")
    category_rule_service.create_rule(keyword="수입전용", match_field="description",
                                      suggested_category_id=category_service.get_income_categories()[0]["id"])
    check("16. 카드에 맞지 않는 추천(수입 카테고리)은 걸러냄",
          "category_id" not in card_service.suggest_for_merchant("수입전용 가맹점"))

    # ---------------- 통장 카드값 결제 맞춰보기 ----------------
    candidates = card_service.find_settlement_candidates(sid)
    check("17. 같은 금액·출금·60일 안의 통장 거래만 후보", [c["id"] for c in candidates] == [bank_settle_id])

    card_service.link_settlement(sid, bank_settle_id)
    bank_tx = transaction_service.get_transaction(bank_settle_id)
    check("18. 연결 시 통장 출금을 비비용출금 / 미지급금으로 변경",
          bank_tx["accounting_type"] == "비비용출금" and bank_tx["account_name"] == "미지급금")
    check("19. 명세서에 연결 정보 표시",
          next(s for s in card_service.list_statements() if s["id"] == sid)["settlement_date"] == "2026-09-15")

    other_sid = card_service.create_statement("다른카드", "2026-08-01", "2026-08-31", AUG_TOTAL,
                                              [dict(use_date="2026-08-01", merchant="x", amount=AUG_TOTAL)])
    cand_other = card_service.find_settlement_candidates(other_sid)
    check("20. 이미 다른 명세서에 연결된 출금은 표시해서 구분", cand_other and cand_other[0]["linked_elsewhere"])
    expect_error("21. 이미 다른 명세서에 연결된 출금에 중복 연결 거부",
                 lambda: card_service.link_settlement(other_sid, bank_settle_id))
    wrong_amount_tx = transaction_service.create_transaction(
        transaction_date="2026-09-15", transaction_type="expense", description="금액 다름", amount=1_000)
    expect_error("22. 금액이 다른 출금과는 연결 거부", lambda: card_service.link_settlement(other_sid, wrong_amount_tx))

    # ---------------- 대시보드 총비용 ----------------
    data = dashboard_service.get_dashboard_data(date(2026, 8, 1), date(2026, 9, 30))
    acct = data["accounting_summary"]
    check("23. 대시보드 총비용 = 통장 비용 + 카드 '비용' 사용내역",
          acct["card_cost"] == AUG_TOTAL and acct["total_cost"] == acct["bank_cost"] + AUG_TOTAL)
    check("24. 연결된 카드값 결제 출금은 통장 비용에서 빠짐(두 번 계산 안 됨)", acct["bank_cost"] == 0)

    # ---------------- 세무사 전달용 Excel ----------------
    wb = load_workbook(io.BytesIO(export_service.export_tax_excel({"start_date": "2026-08-01", "end_date": "2026-08-31"})))
    check("25. 세무사 Excel에 '법인카드 사용내역' 시트 추가", "법인카드 사용내역" in wb.sheetnames)
    card_ws = wb["법인카드 사용내역"]
    amounts = [row[2] for row in card_ws.iter_rows(min_row=2, values_only=True)]
    check("26. 카드 시트에 기간 내 사용내역만 금액 숫자로 포함", sorted(amounts) == sorted([27_000, 14_600, 16_700, 28_457, AUG_TOTAL]))
    wb_no_card = load_workbook(io.BytesIO(export_service.export_tax_excel({"start_date": "2025-01-01", "end_date": "2025-01-31"})))
    check("27. 카드 사용내역이 없는 기간이면 카드 시트를 만들지 않음", wb_no_card.sheetnames == ["세무사 전달용"])

    # ---------------- 해제 / 삭제 ----------------
    card_service.unlink_settlement(sid)
    check("28. 연결 해제", card_service.get_statement(sid)["settlement_transaction_id"] is None)
    card_service.link_settlement(sid, bank_settle_id, reclassify=False)
    result = transaction_service.delete_transactions([bank_settle_id])
    check("29. 명세서와 연결된 통장 출금도 삭제 가능(연결 자동 해제)",
          result["deleted"] == 1 and card_service.get_statement(sid)["settlement_transaction_id"] is None)
    card_service.delete_statement(sid)
    check("30. 명세서 삭제 시 사용내역도 함께 삭제",
          card_service.get_statement(sid) is None and card_service.get_lines(statement_id=sid) == [])

    real_db_path = Path(__file__).resolve().parent.parent / "data" / "finance.db"
    check("31. 실제 data/finance.db는 생성/변경되지 않음", database.DB_PATH == TEMP_DB_PATH and not real_db_path.exists())

    print()
    total = len(results)
    passed = sum(1 for _, ok in results if ok)
    print(f"결과: {passed}/{total} 통과")

    TEMP_DB_PATH.unlink(missing_ok=True)
    shutil.rmtree(BACKUP_DIR, ignore_errors=True)
    return passed == total


if __name__ == "__main__":
    success = main()
    sys.exit(0 if success else 1)
