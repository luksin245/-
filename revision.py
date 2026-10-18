"""신구대조표: 개정안 초안 작성, 개정본 비교, Word 내보내기."""

import difflib
import html
import io
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from checks import AGE, EQUAL, LABOR, OK, RETIRE, norm

NEW, CHANGE, DELETE = "신설", "개정", "삭제"
CIRCLED = [chr(c) for c in range(0x2460, 0x2474)]  # ①~⑳

LAW_SHORT = {LABOR: "근로기준법", EQUAL: "남녀고용평등법", RETIRE: "퇴직급여법", AGE: "고령자고용법"}

# 필수 기재사항이 빠졌을 때 넣을 신설 조항 초안 (고용노동부 표준취업규칙(2026. 2.) 참고)
TEMPLATES = {
    "목적": "제○조(목적) 이 규칙은 회사 직원의 채용·복무 및 근로조건 등에 관한 사항을 정함을 목적으로 한다.",
    "적용범위": "제○조(적용범위) ① 이 규칙은 회사에 근무하는 직원에게 적용한다.\n② 직원의 복무 및 근로조건에 관하여 법령, 단체협약 또는 이 규칙 외의 다른 회사 규정에 별도로 정한 경우를 제외하고는 이 규칙이 정하는 바에 따른다.",
    "연장·야간·휴일근로": "제○조(연장·야간 및 휴일근로) ① 연장근로는 1주간 12시간을 한도로 직원의 동의하에 실시할 수 있다. 다만, 18세 미만 직원은 1일 1시간, 1주 5시간을 한도로 하고, 산후 1년이 지나지 아니한 여성 직원은 1일 2시간, 1주 6시간, 1년 150시간을 한도로 하며, 임신 중인 여성 직원은 연장근로를 시키지 아니한다.\n② 연장·야간근로에 대해서는 통상임금의 100분의 50을 가산하여 지급한다. 단, 휴일근로 중 8시간을 초과한 근로에 대하여는 통상임금의 100분의 100을 가산하여 지급한다.",
    "유급휴일(주휴일)": "제○조(유급휴일) ① 1주 동안 소정근로일을 개근한 직원에게는 매주 ○요일을 유급주휴일로 준다.\n② 근로자의 날(5월 1일)은 유급휴일로 한다.",
    "공휴일 유급휴일": "제○조(공휴일) 관공서의 공휴일에 관한 규정에 따른 공휴일(일요일 제외) 및 대체공휴일은 유급휴일로 한다. 다만, 근로자대표와 서면 합의한 경우 특정한 근로일로 대체할 수 있다.",
    "연차유급휴가": "제○조(연차유급휴가) ① 1년간 80퍼센트 이상 출근한 직원에게는 15일의 유급휴가를 준다.\n② 계속하여 근로한 기간이 1년 미만인 직원 또는 1년간 80퍼센트 미만 출근한 직원에게 1개월 개근 시 1일의 유급휴가를 준다.\n③ 3년 이상 근속한 직원에 대하여는 최초 1년을 초과하는 계속 근로연수 매 2년에 대하여 1일을 가산한 유급휴가를 주며, 가산휴가를 포함한 총 휴가일수는 25일을 한도로 한다.",
    "임금 구성항목": "제○조(임금의 구성항목) 직원의 임금은 기본급 및 ○○수당과 연장·야간·휴일근로수당 등 법정수당으로 구성한다.",
    "임금명세서 교부": "제○조(임금명세서) 회사는 직원에게 임금을 지급하는 때에 임금의 총액, 구성항목, 계산방법, 공제내역, 지급일 등을 기재한 임금명세서를 서면(전자문서를 포함한다)으로 교부한다.",
    "퇴직 사유·퇴직일": "제○조(퇴직 및 퇴직일) ① 직원이 다음 각 호의 어느 하나에 해당하는 경우에는 퇴직으로 한다.\n1. 본인이 퇴직을 원하는 경우\n2. 사망한 경우\n3. 정년에 도달한 경우\n4. 근로계약기간이 만료된 경우\n5. 해고가 결정된 경우\n② 퇴직일은 사직원상 퇴직일, 사망한 날, 정년에 도달한 날, 근로계약기간이 만료된 날, 해고일로 한다.",
    "정년": "제○조(정년) 정년은 만 60세에 도달한 날로 한다.",
    "유산·사산휴가": "제○조(유산·사산휴가) 임신 중인 여성 직원이 유산 또는 사산한 경우로서 해당 직원이 청구하는 경우에는 다음 각 호에 따른 휴가를 준다. 다만, 모자보건법에서 허용되지 않는 인공 임신중절 수술에 따른 경우는 제외한다.\n1. 임신기간이 15주 이내인 경우: 유산 또는 사산한 날부터 10일까지\n2. 임신기간이 16주 이상 21주 이내인 경우: 유산 또는 사산한 날부터 30일까지\n3. 임신기간이 22주 이상 27주 이내인 경우: 유산 또는 사산한 날부터 60일까지\n4. 임신기간이 28주 이상인 경우: 유산 또는 사산한 날부터 90일까지",
    "임신기 근로시간 단축": "제○조(임신기 근로시간 단축) ① 회사는 임신 후 12주 이내 또는 32주 이후에 있는 여성 직원(유산, 조산 등 위험이 있는 경우에는 임신 전 기간)이 1일 2시간의 근로시간 단축을 신청하는 경우 이를 허용한다. 다만, 1일 근로시간이 8시간 미만인 직원에게는 1일 근로시간이 6시간이 되도록 근로시간 단축을 허용할 수 있다.\n② 회사는 제1항에 따른 근로시간 단축을 이유로 해당 직원의 임금을 삭감하지 아니한다.",
    "임신 중 업무 시작·종료 시각 변경": "제○조(임신 중 출퇴근 시각 변경) 회사는 임신 중인 여성 직원이 1일 소정근로시간을 유지하면서 업무의 시작 및 종료 시각의 변경을 신청하는 경우 이를 허용한다. 다만, 정상적인 사업 운영에 중대한 지장을 초래하는 경우 등 근로기준법 시행령으로 정하는 경우에는 그러하지 아니하다.",
    "태아검진 시간": "제○조(태아검진 시간의 허용) ① 회사는 임신한 여성 직원이 모자보건법에 따른 임산부 정기건강진단을 받는 데 필요한 시간을 청구하는 경우 이를 허용한다.\n② 회사는 제1항에 따른 건강진단 시간을 이유로 직원의 임금을 삭감하지 아니한다.",
    "임산부·18세 미만 야간·휴일근로 제한": "제○조(야간 및 휴일근로의 제한) ① 18세 이상의 여성 직원을 오후 10시부터 오전 6시까지 근로하게 하거나 휴일에 근로하게 하려면 그 직원의 동의를 받아야 한다.\n② 임산부와 18세 미만인 직원은 오후 10시부터 오전 6시까지의 시간 및 휴일에 근로시키지 않는 것을 원칙으로 한다. 다만, 18세 미만자의 동의가 있는 경우, 산후 1년이 지나지 아니한 여성의 동의가 있는 경우, 임신 중의 여성이 명시적으로 청구하는 경우로서 근로자대표와 성실히 협의한 후 고용노동부장관의 인가를 받은 경우에는 그러하지 아니하다.",
    "생리휴가": "제○조(생리휴가) 회사는 여성 직원이 청구하는 경우 월 1일의 무급 생리휴가를 준다.",
    "육아시간(수유시간)": "제○조(육아시간) 생후 1년 미만의 자녀가 있는 여성 직원이 청구하는 경우 휴게시간 외에 1일 2회 각각 30분 이상의 유급 수유시간을 준다.",
    "난임치료휴가": "제○조(난임치료휴가) ① 회사는 직원이 인공수정 또는 체외수정 등 난임치료를 받기 위하여 휴가를 청구하는 경우 연간 6일 이내의 휴가를 주며, 이 경우 최초 2일은 유급으로 한다. 다만, 직원이 청구한 시기에 휴가를 주는 것이 정상적인 사업 운영에 중대한 지장을 초래하는 경우에는 직원과 협의하여 그 시기를 변경할 수 있다.\n② 회사는 난임치료휴가를 이유로 해고, 징계 등 불리한 처우를 하지 아니하며, 휴가 청구 업무 처리 과정에서 알게 된 사실을 직원의 의사에 반하여 누설하지 아니한다.",
    "가족돌봄휴직": "제○조(가족돌봄휴직) 회사는 직원이 조부모, 부모, 배우자, 배우자의 부모, 자녀 또는 손자녀의 질병, 사고, 노령으로 인하여 그 가족을 돌보기 위하여 휴직을 신청하는 경우 연간 90일 이내(1회 30일 이상)에서 이를 허용한다.",
    "가족돌봄휴가": "제○조(가족돌봄휴가) ① 회사는 직원이 가족의 질병, 사고, 노령 또는 자녀의 양육으로 인하여 긴급하게 그 가족을 돌보기 위하여 휴가를 신청하는 경우 이를 허용한다. 다만, 직원이 청구한 시기에 휴가를 주는 것이 정상적인 사업 운영에 중대한 지장을 초래하는 경우에는 직원과 협의하여 그 시기를 변경할 수 있다.\n② 가족돌봄휴가 기간은 연간 최장 10일로 하며, 가족돌봄휴직 기간에 포함된다.",
    "가족돌봄 등을 위한 근로시간 단축": "제○조(가족돌봄 등을 위한 근로시간 단축) ① 회사는 직원이 다음 각 호의 어느 하나에 해당하는 사유로 근로시간의 단축을 신청하는 경우 이를 허용한다. 다만, 대체인력 채용이 불가능한 경우 등 남녀고용평등법 시행령으로 정하는 경우에는 그러하지 아니하다.\n1. 가족의 질병, 사고, 노령으로 인하여 그 가족을 돌보기 위한 경우\n2. 직원 자신의 질병이나 사고로 인한 부상 등의 사유로 자신의 건강을 돌보기 위한 경우\n3. 55세 이상의 직원이 은퇴를 준비하기 위한 경우\n4. 직원의 학업을 위한 경우\n② 단축 후 근로시간은 주당 15시간 이상 30시간 이하로 하며, 단축 기간은 1년 이내로 한다. 다만, 제1항제1호부터 제3호까지의 사유로 단축한 직원은 합리적 이유가 있는 경우 2년의 범위에서 연장할 수 있다.\n③ 회사는 근로시간 단축을 이유로 해당 직원에게 해고나 그 밖의 불리한 처우를 하지 아니한다.",
    "안전보건교육": "제○조(안전보건교육) 회사는 직원의 산업재해 예방을 위하여 정기교육, 채용 시 교육, 작업내용 변경 시 교육 등 산업안전보건법령에 따른 안전·보건교육을 실시하며, 직원은 이 교육에 성실하게 참여하여야 한다.",
    "건강진단": "제○조(건강진단) ① 회사는 산업안전보건법령에 따라 매년 1회(사무직은 2년에 1회) 일반건강진단을 실시하고, 필요한 경우 특수·배치전·수시·임시건강진단을 실시한다.\n② 직원은 회사가 실시하는 건강진단을 성실히 받아야 한다.",
    "산업안전보건법 준수": "제○조(산업안전보건법 준수) 회사와 직원은 이 규칙에서 정하지 아니한 사항에 대하여는 산업안전보건법령에서 정하는 바에 따른다.",
    "괴롭힘 금지": "제○조(직장 내 괴롭힘 행위의 금지) ① 직장 내 괴롭힘 행위란 사용자 또는 직원이 직장에서의 지위 또는 관계 등의 우위를 이용하여 업무상 적정범위를 넘어 다른 직원에게 신체적·정신적 고통을 주거나 근무환경을 악화시키는 행위를 말한다.\n② 누구든지 직장 내 괴롭힘 행위를 하여서는 아니 된다.",
    "예방교육": "제○조(직장 내 괴롭힘 예방교육) ① 회사는 직장 내 괴롭힘 예방을 위한 교육을 연 1회 이상 실시하며, 사용자와 직원은 이 교육을 받아야 한다.\n② 회사는 직장 내 괴롭힘 예방교육의 주요 내용을 항상 게시하거나 열람할 수 있도록 한다.",
    "신고·접수": "제○조(직장 내 괴롭힘 사건의 접수) ① 회사는 인사부서에 직장 내 괴롭힘 예방·대응 업무를 담당하는 직원을 1명 이상 둔다.\n② 누구든지 직장 내 괴롭힘 발생 사실을 알게 된 경우 그 사실을 사용자 또는 예방·대응 담당자에게 신고할 수 있다.",
    "조사": "제○조(직장 내 괴롭힘 사건의 조사) 회사는 직장 내 괴롭힘 신고를 접수하거나 발생 사실을 인지한 경우에는 지체 없이 그 사실 확인을 위하여 객관적으로 조사하고, 조사가 끝나면 사건 당사자에게 처리 결과와 판단 근거를 통지한다.",
    "피해자 보호": "제○조(피해직원 등의 보호) ① 회사는 조사기간 동안 피해를 입은 직원 또는 피해를 입었다고 주장하는 직원을 보호하기 위하여 필요한 경우 근무장소의 변경, 유급휴가 명령 등 적절한 조치를 한다. 이 경우 피해직원 등의 의사에 반하는 조치를 하지 아니한다.\n② 회사는 직장 내 괴롭힘이 인정된 경우 피해직원이 요청하면 근무장소의 변경, 배치전환, 유급휴가의 명령 등 적절한 조치를 한다.",
    "행위자 조치": "제○조(직장 내 괴롭힘 행위자에 대한 조치) 회사는 직장 내 괴롭힘이 인정된 경우 지체 없이 행위자에 대하여 징계, 근무장소의 변경 등 필요한 조치를 하며, 조치 전에 피해직원의 의견을 듣는다.",
    "불리한 처우 금지": "제○조(불리한 처우 금지) 회사는 직장 내 괴롭힘 발생 사실을 신고한 직원 및 피해직원 등에게 해고나 그 밖의 불리한 처우를 하지 아니한다.",
    "비밀 누설 금지": "제○조(비밀유지) 직장 내 괴롭힘 발생 사실을 조사한 사람, 조사 내용을 보고받은 사람 및 그 밖에 조사 과정에 참여한 사람은 조사 과정에서 알게 된 비밀을 피해직원 등의 의사에 반하여 다른 사람에게 누설하여서는 아니 된다. 다만, 조사와 관련된 내용을 사용자에게 보고하거나 관계 기관의 요청에 따라 필요한 정보를 제공하는 경우는 제외한다.",
    "징계 사유·종류": "제○조(징계의 종류) 징계는 견책, 감급, 정직, 해고로 한다. 감급은 1회의 금액이 평균임금의 1일분의 2분의 1을, 총액이 1임금지급기의 임금 총액의 10분의 1을 초과하지 못한다.",
    "징계 절차(심의·소명)": "제○조(징계심의) ① 징계는 인사위원회의 심의·의결을 거쳐 결정하며, 위원회는 회의 7일 전까지 징계대상 직원에게 회의 일시, 장소 등을 서면으로 통보한다.\n② 위원회는 의결 전에 징계대상 직원에게 소명할 기회를 준다.",
    "징계 결과 통보": "제○조(징계결과 통보) 회사는 징계 결과를 징계처분 사유와 함께 해당 직원에게 서면으로 통보한다.",
    "시업·종업 시각": "제○조(근로시간) ① 1주간의 근로시간은 휴게시간을 제외하고 40시간, 1일의 근로시간은 휴게시간을 제외하고 8시간으로 한다.\n② 시업시각은 ○○시, 종업시각은 ○○시로 한다.",
    "휴게시간": "제○조(휴게) 회사는 근로시간이 4시간인 경우에는 30분 이상, 8시간인 경우에는 1시간 이상의 휴게시간을 근로시간 도중에 준다.",
    "결정·계산·지급 방법": "제○조(임금의 계산 및 지급방법) 임금은 월 단위로 계산하여 통화로 직접 직원에게 그 전액을 지급한다. 다만, 직원의 동의를 받아 직원 명의의 예금계좌로 지급할 수 있다.",
    "산정기간·지급시기": "제○조(임금의 산정기간 및 지급일) 임금은 매월 1일부터 말일까지를 산정기간으로 하여 다음 달 ○일에 지급한다. 지급일이 휴일인 경우에는 그 전날 지급한다.",
    "승급": "제○조(승급) 승급은 필요시 실시하며, 그 기준과 방법은 회사가 별도로 정한다.",
    "퇴직급여": "제○조(퇴직급여) 회사는 계속근로기간이 1년 이상인 직원이 퇴직하는 경우 근로자퇴직급여 보장법에 따라 계속근로기간 1년에 대하여 30일분 이상의 평균임금을 퇴직금으로 지급하거나 퇴직연금제도를 운영한다.",
    "최저임금": "제○조(최저임금) 회사는 최저임금법에 따른 최저임금액 이상의 임금을 지급한다.",
    "출산전후휴가": "제○조(출산전후휴가) ① 회사는 임신 중인 여성 직원에게 출산 전과 출산 후를 통하여 90일(미숙아를 출산한 경우에는 100일, 한 번에 둘 이상 자녀를 임신한 경우에는 120일)의 출산전후휴가를 준다. 이 경우 출산 후에 45일(한 번에 둘 이상 자녀를 임신한 경우에는 60일) 이상이 되도록 한다.\n② 휴가 기간 중 최초 60일(한 번에 둘 이상 자녀를 임신한 경우에는 75일)은 유급으로 한다.",
    "배우자 출산휴가": "제○조(배우자 출산휴가) ① 회사는 직원이 배우자의 출산을 이유로 휴가를 고지하는 경우 근무일 기준으로 20일의 유급휴가를 준다.\n② 제1항에 따른 휴가는 배우자가 출산한 날부터 120일이 지나면 사용할 수 없으며, 3회에 한정하여 나누어 사용할 수 있다.",
    "육아휴직": "제○조(육아휴직) ① 회사는 임신 중인 여성 직원이 모성을 보호하거나 만 8세 이하 또는 초등학교 2학년 이하의 자녀(입양한 자녀를 포함한다)를 양육하기 위하여 휴직을 신청하는 경우 이를 허용한다.\n② 육아휴직 기간은 1년 이내로 한다. 다만, 같은 자녀에 대하여 부모가 각각 3개월 이상 육아휴직을 사용한 경우의 부 또는 모, 한부모, 고용노동부령으로 정하는 장애아동의 부 또는 모는 6개월 이내에서 추가로 사용할 수 있다.\n③ 직원은 육아휴직을 3회에 한정하여 나누어 사용할 수 있다.",
    "육아기 근로시간 단축": "제○조(육아기 근로시간 단축) ① 회사는 직원이 만 12세 이하 또는 초등학교 6학년 이하의 자녀를 양육하기 위하여 근로시간의 단축을 신청하는 경우 이를 허용한다. 다만, 남녀고용평등법 시행령으로 정하는 경우에는 그러하지 아니하다.\n② 단축 후 근로시간은 주당 15시간 이상 35시간 이하로 하며, 단축 기간은 1년 이내로 하되 육아휴직 기간 중 사용하지 아니한 기간이 있으면 그 기간의 두 배를 가산한 기간 이내로 한다.\n③ 직원은 육아기 근로시간 단축을 나누어 사용할 수 있으며, 1회의 기간은 1개월 이상이어야 한다.",
    "사업장 환경 개선": "제○조(사업장 환경의 개선) 회사는 직원의 성별·연령 또는 신체적 조건 등의 특성을 고려하여 사업장 환경을 개선하도록 노력한다.",
    "업무상 재해": "제○조(재해보상) 직원이 업무상 부상을 당하거나 질병에 걸린 경우에는 산업재해보상보험법에 따라 보상을 받는다.",
    "업무 외 재해부조": "제○조(업무 외 재해부조) 직원이 업무 외의 사유로 부상·질병 또는 사망한 경우 회사는 별도로 정하는 바에 따라 부조할 수 있다.",
    "표창": "제○조(표창) 회사는 업무 성과가 우수하거나 회사 발전에 공로가 있는 직원을 표창할 수 있다.",
}


JOSA = [("을", "를"), ("이", "가"), ("은", "는"), ("과", "와"), ("으로", "로")]


def _has_batchim(word):
    """마지막 글자에 받침이 있는지. 숫자는 읽는 소리로 판단한다(10→십, 2→이)."""
    last = word.rstrip()[-1:]
    if not last:
        return False
    if last.isdigit():
        return last == "0" or last not in "2459"
    if "가" <= last <= "힣":
        return (ord(last) - 0xAC00) % 28 != 0
    return False


def _ends_with_rieul(word):
    last = word.rstrip()[-1:]
    if last.isdigit():
        return last in "178"
    return "가" <= last <= "힣" and (ord(last) - 0xAC00) % 28 == 8


def _sub_once(pattern, repl, text):
    """pattern을 repl로 한 번 바꾸고, 바로 뒤 조사(을/를 등)를 repl에 맞게 고친다."""
    m = re.search(pattern, text)
    if not m:
        return text
    rest = text[m.end():]
    for with_b, without_b in JOSA:
        for josa in (with_b, without_b):
            if rest.startswith(josa) and not rest.startswith(with_b + with_b[-1:]):
                if with_b == "으로":
                    right = "로" if (not _has_batchim(repl) or _ends_with_rieul(repl)) else "으로"
                else:
                    right = with_b if _has_batchim(repl) else without_b
                rest = right + rest[len(josa):]
                break
        else:
            continue
        break
    return text[:m.start()] + repl + rest


def _row(kind, label, current, revised, reason, position=None, basis=()):
    """position: 신설 조항을 넣을 자리(이 조항 다음). None이면 본칙 맨 끝. basis: 근거 법령 [(법령명, 조문)]"""
    return {"구분": kind, "조항": label, "현행": current, "개정안": revised, "사유": reason, "위치": position,
            "근거": list(basis)}


def _basis_text(basis):
    return ", ".join(f"{LAW_SHORT.get(law, law)} {label}" for law, label in basis)


def _append_item(text, sentence):
    """조항 끝에 문장을 덧붙인다. 항(①②…)으로 나뉜 조항이면 다음 항 번호를 붙인다."""
    used = [i for i, c in enumerate(CIRCLED) if c in text]
    if used and used[-1] + 1 < len(CIRCLED):
        sentence = f"{CIRCLED[used[-1] + 1]} {sentence}"
    return f"{text}\n{sentence}"


def _drop_paragraph(text, start):
    """start로 시작하는 항(① …과 그 아래 호)을 지우고, 뒤의 항 번호(③→②)를 다시 매긴다."""
    lines = text.split("\n")
    out, skipping = [], False
    for line in lines:
        head = line.strip()
        if head.startswith(start):
            skipping = True
            continue
        if skipping and head[:1] in CIRCLED:
            skipping = False
        if not skipping:
            out.append(line)
    n = 0
    for i, line in enumerate(out):
        head = line.lstrip()
        if head[:1] in CIRCLED:
            out[i] = line[:len(line) - len(head)] + CIRCLED[n] + head[1:]
            n += 1
    return "\n".join(out)


def _required_basis(no):
    return f"근로기준법 제93조 제{no}호" if no[0].isdigit() else f"표준취업규칙 {no}"


def draft_rows(articles, required, findings):
    """점검 결과로 신구대조표 초안을 만든다."""
    by_label = {a["조항"]: a for a in articles}
    changed = {}  # 조항 → [개정안, 사유 목록, {원래 문장: 고친 문장}]
    new_rows = []

    for f in findings:
        fix = f.get("fix")
        if not fix or f["상태"] == OK:
            continue
        reason = f"{f['주제']}({_basis_text(f['근거'])})" if f["근거"] else f["주제"]
        if "new" in fix:
            near = [a["조항"] for a in articles if not a["부칙"] and "괴롭힘" in a["제목"] + a["내용"]]
            new_rows.append(_row(NEW, "", "", fix["new"], reason, near[-1] if near else None, f["근거"]))
            continue
        article = by_label.get(f["조항"])
        if not article:
            continue
        entry = changed.setdefault(article["조항"], [article["내용"], [], {}, []])
        if "drop_para" in fix:
            dropped = _drop_paragraph(entry[0], fix["drop_para"])
            if dropped == entry[0]:
                continue
            entry[0] = dropped
        elif "pattern" in fix or "subs" in fix:
            # 같은 문장에 고칠 곳이 여럿이면 앞에서 고친 문장에 이어서 고친다
            current = entry[2].get(f["발췌"], f["발췌"])
            fixed = current
            for pattern, repl in fix.get("subs", [(fix.get("pattern"), fix.get("repl"))]):
                fixed = _sub_once(pattern, repl, fixed)
            if fixed == current or current not in entry[0]:
                continue
            entry[0] = entry[0].replace(current, fixed, 1)
            entry[2][f["발췌"]] = fixed
        else:
            entry[0] = _append_item(entry[0], fix["add"])
        entry[1].append(reason)
        entry[3].extend(f["근거"])

    rows = [_row(CHANGE, label, by_label[label]["내용"], strip_internal_refs(text), "\n".join(reasons), basis=basis)
            for label, (text, reasons, _, basis) in changed.items()]
    rows.sort(key=lambda r: articles.index(by_label[r["조항"]]))

    order = {a["조항"]: i for i, a in enumerate(articles)}
    chapter_of = {a["조항"]: a["장"] for a in articles}
    for item in required:
        # 같은 항목의 다른 세부가 가장 많이 들어 있는 조항의 장을 찾고, 그 장 안의 관련 조항 중 마지막 다음에 넣는다
        # (예: 가족돌봄휴가 → 모성보호 장의 마지막 관련 조항 뒤, 승급 → 임금 장의 관련 조항 뒤)
        found = [label for d in item["세부"] for label in d["조항"] if label in order]
        position = None
        if found:
            best = max(set(found), key=lambda x: (found.count(x), order[x]))
            same = [x for x in found if chapter_of[x] == chapter_of[best]]
            position = max(same, key=order.get)
        for d in item["세부"]:
            if not d["조항"] and d["세부"] in TEMPLATES:
                new_rows.append(_row(NEW, "", "", TEMPLATES[d["세부"]],
                                     f"필수 기재사항 누락: {d['세부']}({_required_basis(item['번호'])})", position))
    return rows + new_rows


def _body(text):
    """조항 번호를 뺀 본문(공백 제거)."""
    return norm(re.sub(r"^\s*제\s*\d+\s*조(?:\s*의\s*\d+)?", "", text))


def compare_rows(old_articles, new_articles):
    """현행 취업규칙과 개정본을 비교한다.

    조항 번호가 바뀌어도 같은 조항으로 볼 수 있도록 제목(본칙/부칙 구분 포함)으로 먼저 맞추고,
    제목으로 못 찾으면 조항 번호로 맞춘다.
    """
    def title_key(a):
        return (a["부칙"], norm(a["제목"])) if a["제목"] else None

    titles = [title_key(a) for a in old_articles]
    by_title = {t: a for t, a in zip(titles, old_articles) if t and titles.count(t) == 1}
    by_label = {a["조항"]: a for a in old_articles}
    used = set()
    rows = []
    for a in new_articles:
        before = by_title.get(title_key(a))
        if before is None or before["조항"] in used:
            before = by_label.get(a["조항"])
        if before is None or before["조항"] in used:
            rows.append(_row(NEW, a["조항"], "", a["내용"], ""))
            continue
        used.add(before["조항"])
        # 조항 번호만 바뀐 경우(제7조 → 제6조)는 개정으로 보지 않는다
        if _body(before["내용"]) != _body(a["내용"]):
            label = a["조항"] if a["조항"] == before["조항"] else f"{before['조항']} → {a['조항']}"
            rows.append(_row(CHANGE, label, before["내용"], a["내용"], ""))
    for a in old_articles:
        if a["조항"] not in used:
            rows.append(_row(DELETE, a["조항"], a["내용"], "", ""))

    def order(row):
        m = re.match(r"(부칙 )?제(\d+)조(?:의(\d+))?", row["조항"])
        return (bool(m.group(1)), int(m.group(2)), int(m.group(3) or 0)) if m else (True, 9999, 0)

    return sorted(rows, key=order)


# ---------------------------------------------------------------------------
# 변경 부분 표시
# ---------------------------------------------------------------------------
def _tokens(text):
    return re.findall(r"\s+|[^\s]+", text)


def diff_segments(current, revised):
    """(현행 조각, 개정안 조각)을 돌려준다. 조각은 (글자, 바뀐 부분인지)."""
    a, b = _tokens(current), _tokens(revised)
    left, right = [], []
    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
        changed = op != "equal"
        if i2 > i1:
            left.append(("".join(a[i1:i2]), changed))
        if j2 > j1:
            right.append(("".join(b[j1:j2]), changed))
    return left, right


def diff_html(segments):
    parts = []
    for text, changed in segments:
        t = html.escape(text).replace("\n", "<br>")
        parts.append(f"<u><b style='color:#c0392b'>{t}</b></u>" if changed and text.strip() else t)
    return "".join(parts)


# ---------------------------------------------------------------------------
# Word 내보내기 (A4 가로, 제목 상자 + 현행/개정 안/비고 3칸 표)
# ---------------------------------------------------------------------------
FONTS = ("함초롬바탕", "바탕", "맑은 고딕")
HEADING = re.compile(r"^(제\s*[\d○]+\s*조(?:\s*의\s*\d+)?)\s*([(（【\[][^)）】\]]*[)）】\]])?\s*(.*)$", re.S)
ITEM_INDENT = [  # (문단 시작 패턴, 왼쪽 들여쓰기 cm) — 둘째 줄부터 글머리 뒤에 맞춘다
    (re.compile(rf"^[{''.join(CIRCLED)}]"), 0.45),
    (re.compile(r"^\d+\.\s"), 0.9),
    (re.compile(r"^[가나다라마바사아자차카타파하]\.\s"), 1.35),
]


def _layout(text):
    """첫 줄의 '제17조(변경사항 신고) 본문'을 '제17조 (변경사항 신고)' 줄과 본문 줄로 나눈다."""
    lines = [line.strip() for line in text.strip().split("\n")]
    m = HEADING.match(lines[0]) if lines and lines[0] else None
    if not m:
        return "\n".join(lines), False
    head = re.sub(r"\s+", "", m.group(1)) + (f" {m.group(2)}" if m.group(2) else "")
    rest = [m.group(3).strip()] if m.group(3).strip() else []
    return "\n".join([head] + rest + lines[1:]), True


def _font(run, name, size, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    # 굵게가 아닐 때는 '굵게 해제' 표시를 쓰지 않는다(일부 프로그램이 이를 굵게로 잘못 읽는다)
    run.font.bold = True if bold else None
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:eastAsia", "w:ascii", "w:hAnsi"):
        rfonts.set(qn(attr), name)


def _shade(cell, color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shd)


def _tight(paragraph):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.15


def _indent(paragraph, text):
    for pattern, cm in ITEM_INDENT:
        if pattern.match(text):
            paragraph.paragraph_format.left_indent = Cm(cm)
            paragraph.paragraph_format.first_line_indent = Cm(-0.45)
            return


def _write_cell(cell, segments, has_heading, font, size, highlight):
    """조각들을 문단으로 나눠 쓴다. 첫 문단이 조 제목이면 굵게, 항·호·목은 내어쓰기."""
    paras = "".join(t for t, _ in segments).split("\n")
    p_index, paragraph = 0, cell.paragraphs[0]
    _tight(paragraph)
    _indent(paragraph, paras[0])
    for text, changed in segments:
        for k, piece in enumerate(text.split("\n")):
            if k:
                p_index += 1
                paragraph = cell.add_paragraph()
                _tight(paragraph)
                _indent(paragraph, paras[p_index])
            if not piece:
                continue
            run = paragraph.add_run(piece)
            _font(run, font, size, bold=has_heading and p_index == 0)
            if highlight and changed and piece.strip():
                run.font.underline = True  # 바뀐 부분은 밑줄만 긋는다
    paragraph.paragraph_format.space_after = Pt(10)  # 조항 사이 여백


def _placeholder(cell, text, font, size):
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(p.add_run(text), font, size)


def _title_box(doc, title, font, page_width_cm):
    """그림자가 있는 테두리 상자 안의 제목."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    side = Cm(max(1.0, (page_width_cm - 10.5) / 2))
    p.paragraph_format.left_indent = side
    p.paragraph_format.right_indent = side
    p.paragraph_format.space_after = Pt(14)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:space"), "8")
        el.set(qn("w:color"), "000000")
        el.set(qn("w:shadow"), "1")
        borders.append(el)
    p._p.get_or_add_pPr().append(borders)
    _font(p.add_run(title), font, 17, bold=True)


def to_docx(rows, title="취업규칙 신구대조표", font=FONTS[0], highlight=True, with_reason=True, size=9.5,
            company="", written=""):
    """신구대조표. company·written이 있으면 제목 아래에 회사명과 작성일을 적는다."""
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21.0)
    margin = 1.2
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(margin))
    usable = 29.7 - 2 * margin

    _title_box(doc, title, font, usable)
    info = [f"작성일 : {written}" if written.strip() else "", f"회사명 : {company}" if company.strip() else ""]
    info = [x for x in info if x]
    for i, line in enumerate(info):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _tight(p)
        if i == len(info) - 1:
            p.paragraph_format.space_after = Pt(6)
        _font(p.add_run(line), font, size + 0.5)

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Cm(usable * 0.405), Cm(usable * 0.455), Cm(usable * 0.14))
    for column, width in zip(table.columns, widths):
        column.width = width
    for cell, name, width in zip(table.rows[0].cells, ("현행", "개정 안", "비고"), widths):
        cell.width = width
        _shade(cell, "D9D9D9")
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _tight(cell.paragraphs[0])
        _font(cell.paragraphs[0].add_run(name), font, size)

    for r in rows:
        cells = table.add_row().cells
        for cell, width in zip(cells, widths):
            cell.width = width
        current, cur_head = _layout(r["현행"])
        revised, new_head = _layout(r["개정안"])
        left, right = diff_segments(current, revised)
        if r["구분"] == NEW:
            _placeholder(cells[0], "<신 설>", font, size)
        else:
            _write_cell(cells[0], left, cur_head, font, size, highlight)
        if r["구분"] == DELETE:
            _placeholder(cells[1], "<삭 제>", font, size)
        else:
            _write_cell(cells[1], right, new_head, font, size, highlight)
        if with_reason and r["사유"]:
            _tight(cells[2].paragraphs[0])
            _font(cells[2].paragraphs[0].add_run(r["사유"]), font, size - 1)

    if highlight:
        note = doc.add_paragraph()
        _font(note.add_run("※ 밑줄 친 부분이 바뀐 내용입니다."), font, size - 1)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# 개정 취업규칙 초안 (현행 전체에 신구대조표 내용을 반영한 Word 문서)
# ---------------------------------------------------------------------------
CHANGED_COLOR = RGBColor(0x1F, 0x4E, 0xB4)  # 개정 조항: 파란색
NEW_COLOR = RGBColor(0x2E, 0x7D, 0x32)      # 신설 조항: 초록색


def _heading_line(doc, text, font, size, toc=False):
    p = doc.add_paragraph(style="Heading 1" if toc else None)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    _font(p.add_run(text), font, size, bold=True)


def _write_article(doc, text, font, size, color=None, toc=False):
    laid, has_heading = _layout(text)
    for i, line in enumerate(laid.split("\n")):
        heading = has_heading and i == 0
        p = doc.add_paragraph(style="Heading 2" if toc and heading else None)
        _tight(p)
        p.paragraph_format.line_spacing = 1.3
        _indent(p, line)
        run = p.add_run(line)
        _font(run, font, size, bold=heading)
        if color is not None:
            run.font.color.rgb = color
    p.paragraph_format.space_after = Pt(8)


def _plain_heading_styles(doc, font, size):
    """목차용 제목 스타일(장=제목 1, 조=제목 2)을 본문과 같은 모양(검은색, 같은 글꼴)으로 만든다."""
    for name, pt in (("Heading 1", size + 1), ("Heading 2", size)):
        style = doc.styles[name]
        style.font.name = font
        style.font.size = Pt(pt)
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
        style.paragraph_format.keep_with_next = True


def _field_char(run, kind):
    el = OxmlElement("w:fldChar")
    el.set(qn("w:fldCharType"), kind)
    run._r.append(el)


def _table_of_contents(doc, entries, font, size):
    """목차. Word가 문서를 열 때 필드를 업데이트하면 쪽 번호가 붙은 목차로 바뀌고,
    업데이트하지 않아도 장·조 목록은 그대로 보인다."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    _font(p.add_run("목 차"), font, size + 3, bold=True)
    if not entries:
        return
    paras = []
    for level, text in entries:
        q = doc.add_paragraph()
        _tight(q)
        q.paragraph_format.line_spacing = 1.2
        if level == 2:
            q.paragraph_format.left_indent = Cm(0.6)
        else:
            q.paragraph_format.space_before = Pt(4)
        paras.append((q, level, text))
    first, last = paras[0][0], paras[-1][0]
    begin = first.add_run()
    _field_char(begin, "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    begin._r.append(instr)
    _field_char(begin, "separate")
    for q, level, text in paras:
        _font(q.add_run(text), font, size - 0.5, bold=level == 1)
    _field_char(last.add_run(), "end")
    # 문서를 열 때 Word가 필드(목차 쪽 번호) 업데이트를 묻도록 한다
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings.element.append(update)
    doc.add_page_break()


def _toc_line(text):
    laid, has_heading = _layout(text)
    return laid.split("\n", 1)[0] if has_heading else None


def _article_no(label):
    m = re.match(r"제(\d+)조", label or "")
    return int(m.group(1)) if m else 0


LABEL = re.compile(r"제\s*(\d+)\s*조(?:\s*의\s*(\d+))?")
# 이런 말 바로 뒤의 '제n조'는 다른 법령을 가리키므로 번호를 바꾸지 않는다
# ('이 규칙'·'본 규칙'·'취업규칙'은 이 규칙 자신이므로 제외)
OTHER_LAW = re.compile(r"(법|령|」|』|협약|규약|한다\)|(?<!취업)(?<![이본])(?<![이본]\s)(?:규칙|규정))\s*$")


def _set_label(text, label):
    """조항 첫머리의 '제○조'·'제12조의2'를 label로 바꾼다. 첫머리에 조 번호가 없으면 붙인다."""
    text = text.strip()
    if re.match(r"^제\s*[\d○]+\s*조(?:\s*의\s*\d+)?", text):
        return re.sub(r"^제\s*[\d○]+\s*조(?:\s*의\s*\d+)?", label, text, count=1)
    return f"{label} {text}"


def _remap_refs(text, mapping):
    """본문 속 '제n조'(이 규칙 안의 인용)를 새 번호로 바꾼다. 첫머리 조 번호와 다른 법령 인용은 그대로 둔다."""
    head = re.match(r"^\s*제\s*\d+\s*조(?:\s*의\s*\d+)?", text)
    start = head.end() if head else 0

    def repl(m):
        pos = start + m.start()  # m은 text[start:] 기준 위치
        if OTHER_LAW.search(text[max(0, pos - 15):pos]):
            return m.group(0)
        key = f"제{m.group(1)}조" + (f"의{m.group(2)}" if m.group(2) else "")
        return mapping.get(key, m.group(0))

    return text[:start] + LABEL.sub(repl, text[start:])


# 이 규칙의 다른 조항을 가리키는 표현: "제66조 제1항에 따라", "이 규칙 제67조에 따른", "제17조제4호에 의한"
INTERNAL_REF = re.compile(
    r"(?:(?:이|본)\s*규칙\s*)?"
    r"제\s*\d+\s*조(?:\s*의\s*\d+)?(?:\s*제\s*\d+\s*항)?(?:\s*제\s*\d+\s*호)?"
    r"(?:\s*(?:및|또는|,|ㆍ|·)\s*제\s*\d+\s*(?:조|항|호)(?:\s*의\s*\d+)?)*"
    r"\s*(?:의\s*규정에\s*)?(?:에\s*따라|에\s*따른|에\s*의한|에\s*의하여|에서\s*정한|에서\s*정하는)\s*"
)


def strip_internal_refs(text):
    """조문 안의 '제66조 제1항에 따라' 같은 다른 조항 인용을 지운다 (조 번호가 바뀌어도 어긋나지 않게).
    첫머리 조 번호와 다른 법령 인용(「직업안정법」 제2조의2 제1호에 따른)은 그대로 둔다."""
    head = re.match(r"^\s*제\s*[\d○]+\s*조(?:\s*의\s*\d+)?(?:\s*[(（【\[][^)）】\]]*[)）】\]])?", text)
    start = head.end() if head else 0

    def repl(m):
        pos = start + m.start()
        if OTHER_LAW.search(text[max(0, pos - 15):pos]):
            return m.group(0)
        return ""

    return text[:start] + INTERNAL_REF.sub(repl, text[start:])


def plan_rules(articles, rows, renumber=False):
    """개정 후 본칙의 조항 순서와 번호.

    반환: [{"종류": "현행"|"신설", "조항": 새 번호, "이전": 현행 번호, "행": 신구대조표 행, "내용": 개정 후 문구,
            "장": 장 제목}], 번호 대응표 {현행 번호: 새 번호}
    renumber=False: 기존 번호는 그대로 두고, 신설 조항은 가지번호(제40조의2)로 끼워 넣는다.
    renumber=True: 신설·삭제를 반영해 제1조부터 번호를 다시 매기고, 본문 속 조 인용도 고친다.
    """
    body = [a for a in articles if not a["부칙"]]
    labels = [a["조항"] for a in body]
    changed = {r["조항"]: r for r in rows if r["조항"] and r["구분"] in (CHANGE, DELETE)}
    after, tail = {}, []
    for r in rows:
        if r["구분"] == NEW and not r["조항"]:
            if r.get("위치") in labels:
                after.setdefault(r["위치"], []).append(r)
            else:
                tail.append(r)

    seq = []
    for a in body:
        seq.append({"종류": "현행", "이전": a["조항"], "행": changed.get(a["조항"]), "원문": a["내용"], "장": a["장"]})
        for r in after.get(a["조항"], []):
            seq.append({"종류": "신설", "이전": None, "행": r, "위치": a["조항"], "장": a["장"]})
    for r in tail:
        seq.append({"종류": "신설", "이전": None, "행": r, "위치": None, "장": None})

    mapping = {}
    if renumber:
        seq = [it for it in seq if not (it["행"] and it["행"]["구분"] == DELETE)]
        for n, it in enumerate(seq, 1):
            it["조항"] = f"제{n}조"
            if it["이전"]:
                mapping[it["이전"]] = it["조항"]
    else:
        used = set(labels)
        tail_no = max((_article_no(x) for x in labels), default=0) + 1
        for it in seq:
            if it["이전"]:
                it["조항"] = it["이전"]
            elif it["위치"]:
                base, branch = _article_no(it["위치"]), 2
                while f"제{base}조의{branch}" in used:
                    branch += 1
                it["조항"] = f"제{base}조의{branch}"
                used.add(it["조항"])
            else:
                it["조항"] = f"제{tail_no}조"
                tail_no += 1

    # 신설 조항의 제목 괄호를 현행 규칙의 표기(예: 제1조【목적】)에 맞춘다
    lunar = sum("【" in a["내용"].split("\n", 1)[0] for a in body) > len(body) / 2
    for it in seq:
        row = it["행"]
        if it["종류"] == "신설":
            text = row["개정안"]
            if lunar:
                text = re.sub(r"^(제\s*[\d○]+\s*조(?:\s*의\s*\d+)?)\s*\(([^)]*)\)", r"\1【\2】", text.strip(), count=1)
        elif row and row["구분"] == DELETE:
            text = f"{it['이전']} 삭제"
        elif row:
            text = row["개정안"]
        else:
            text = it["원문"]
        text = _set_label(text, it["조항"])
        it["내용"] = _remap_refs(text, mapping) if mapping else text
    return seq, mapping


def numbered_rows(seq, rows):
    """신구대조표 행에 새 조 번호를 반영해 조항 순서대로 돌려준다 (비교 모드 행은 그대로 뒤에 붙인다)."""
    out, seen = [], set()
    for it in seq:
        row = it["행"]
        if row is None:
            continue
        seen.add(id(row))
        revised = "" if row["구분"] == DELETE else it["내용"]
        out.append({**row, "조항": it["조항"] if row["구분"] == NEW else row["조항"], "개정안": revised})
    return out + [r for r in rows if id(r) not in seen]


def _korean_date(text):
    """'2026. 12. 1.', '2026-12-01', '2026년 12월 1일' → ('2026년 12월 1일', '2026. 12. 1.'). 못 읽으면 빈칸 표시."""
    m = re.search(r"(\d{4})\s*[.년/-]\s*(\d{1,2})\s*[.월/-]\s*(\d{1,2})", text or "")
    if not m:
        return "○○○○년 ○월 ○일", "○○○○. ○. ○."
    y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
    return f"{y}년 {mo}월 {d}일", f"{y}. {mo}. {d}."


def to_rules_docx(articles, seq, title="취업규칙 (개정안)", font=FONTS[0], mark=True, size=10.5, effective="",
                  toc=True):
    """개정 후 취업규칙 전체. mark=True면 검토용(개정 파란색·신설 초록색), False면 최종본. toc=True면 앞에 목차."""
    doc = Document()
    _plain_heading_styles(doc, font, size)
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    for side, cm in (("left_margin", 2.5), ("right_margin", 2.5), ("top_margin", 2.0), ("bottom_margin", 2.0)):
        setattr(section, side, Cm(cm))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    _font(p.add_run(title), font, 18, bold=True)
    if mark:
        note = doc.add_paragraph()
        _font(note.add_run("※ 파란색은 개정한 조항, 초록색은 신설한 조항입니다. (검토용)"), font, size - 1.5)

    addenda = [a for a in articles if a["부칙"]]
    long_date, short_date = _korean_date(effective)
    new_addendum = f"부 칙 ({short_date})"

    if toc:
        entries, chapter = [], None
        for it in seq:
            if it["장"] and it["장"] != chapter:
                chapter = it["장"]
                entries.append((1, chapter))
            line = _toc_line(it["내용"])
            if line:
                entries.append((2, line))
        if addenda:
            entries.append((1, "부 칙"))
            entries += [(2, x) for x in (_toc_line(a["내용"]) for a in addenda) if x]
        entries.append((1, new_addendum))
        _table_of_contents(doc, entries, font, size)

    chapter = None
    for it in seq:
        if it["장"] and it["장"] != chapter:
            chapter = it["장"]
            _heading_line(doc, chapter, font, size + 1, toc)
        color = None
        if mark and it["종류"] == "신설":
            color = NEW_COLOR
        elif mark and it["행"]:
            color = CHANGED_COLOR
        _write_article(doc, it["내용"], font, size, color, toc)

    if addenda:
        _heading_line(doc, "부 칙", font, size + 1, toc)
        for a in addenda:
            _write_article(doc, a["내용"], font, size, toc=toc)
    _heading_line(doc, new_addendum, font, size + 1, toc)
    _write_article(doc, f"이 규칙은 {long_date}부터 시행한다.", font, size, NEW_COLOR if mark else None)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
